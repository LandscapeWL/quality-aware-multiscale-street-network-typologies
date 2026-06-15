#!/usr/bin/env python3
"""Step13: export manuscript tables, appendices, provenance, and material lists.

This step is deliberately export-only. It reads stable outputs from previous
steps and writes paper tables, appendix workbooks, material inventories, input
acceptance reports, QC records, execution logs, reproducibility metadata, and a
README under ``data/13_export_paper_tables_appendices``.

It does not remodel, rewrite upstream data, modify Step12 figures, or update
the manuscript document. Table 1 remains manuscript-maintained and is only
documented in the material list and README.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.metadata as importlib_metadata
import json
import math
import os
import platform
import re
import subprocess
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


STEP_NAME = "13_export_paper_tables_appendices"
STEP01_NAME = "01_city_boundaries_and_sample_inventory"
STEP06_NAME = "06_clean_networks_compute_morphology_metrics"
STEP07_NAME = "07_generate_quality_scores_and_type_confidence"
STEP08_NAME = "08_generate_scale_signatures_and_city_profiles"
STEP09_NAME = "09_cluster_morphotypes"
STEP10_NAME = "10_calculate_accessibility_and_detour_validation_metrics"
STEP11_NAME = "11_statistical_models_and_robustness_checks"
STEP12_NAME = "12_draw_paper_figures"

MORPHOTYPES = [f"MT{i:02d}" for i in range(1, 11)]
FACILITY_CATEGORIES = ["grocery", "healthcare", "park", "school", "transit"]

MORPHOTYPE_LABELS_EN = {
    "MT01": "Dense fine-grain core",
    "MT02": "Dense ordered grid",
    "MT03": "Mid-density fine-grain",
    "MT04": "Mid-density cul-de-sac blocks",
    "MT05": "Mid-density long-grid corridors",
    "MT06": "Mid-density mixed fabric",
    "MT07": "Low-density curvilinear network",
    "MT08": "Low-density cul-de-sac fabric",
    "MT09": "Sparse peripheral cul-de-sac",
    "MT10": "Sparse curvilinear superblocks",
}

RISK_NOTES = [
    "Model results are explanatory associations, not causal effects.",
    "Fig8 detour panel p10/p90 intervals are distribution intervals across city-type groups, not confidence intervals.",
    "Step12 Fig9 and Step11 model_marginal_effects.csv are observed weighted group means / group marginal summary, not adjusted marginal effects.",
    "Step11 small-group warning is recorded as a non-blocking warning.",
    "Descriptive morphotype names are not unique; Step13 uses MT01-MT10 as the unique type keys.",
    "MT01-MT10 are local morphotypes; representative cities indicate city composition examples, not city-level classes.",
]

DOMINANT_COMPATIBILITY_NOTE = (
    "Step13 city composition outputs retain dominant_morphotype as a compatibility field equal to "
    "population_dominant_morphotype; dominant_population_share and dominant_unit_share are shares "
    "for that population-dominant MT. Use unit_dominant_* for unit-share dominance."
)

MAIN_TABLE_FILES = [
    "table2_main_sample_region_distribution.csv",
    "table3_city_sample_stratification_logic.csv",
    "table4_data_source_inventory.csv",
    "table5_implementation_steps.csv",
    "table6_analysis_modules.csv",
    "table7_validation_model_design.csv",
    "table8_robustness_checks.csv",
    "table9_morphotype_definitions_and_representative_cities.csv",
]

APPENDIX_TABLE_FILES = [
    "city_morphotype_composition_full.csv",
]

APPENDIX_FILES = [
    "appendix_a_metric_dictionary.xlsx",
    "appendix_b_sample_city_inventory.xlsx",
    "appendix_c_quality_score_details.xlsx",
    "appendix_d_clustering_and_type_naming.xlsx",
    "appendix_e_accessibility_detour_validation.xlsx",
    "appendix_f_robustness_checks.xlsx",
]

RECORD_FILES = [
    "paper_materials_manifest.csv",
    "paper_materials_manifest.xlsx",
    "step13_input_acceptance_report.json",
    "step13_input_acceptance_report.md",
    "step13_quality_checks.csv",
    "step13_run_log.json",
    "reproducibility_status_manifest.json",
    "README.md",
]

OUTPUT_FILES = MAIN_TABLE_FILES + APPENDIX_TABLE_FILES + APPENDIX_FILES + RECORD_FILES


@dataclass(frozen=True)
class StepPaths:
    root: Path
    output_dir: Path
    step01_dir: Path
    step06_dir: Path
    step07_dir: Path
    step08_dir: Path
    step09_dir: Path
    step10_dir: Path
    step11_dir: Path
    step12_dir: Path
    docs_dir: Path

    @property
    def city_master(self) -> Path:
        return self.step01_dir / "city_sample" / "city_master.parquet"

    @property
    def city_boundaries(self) -> Path:
        return self.step01_dir / "city_sample" / "city_boundaries.gpkg"

    @property
    def city_morphology_metrics(self) -> Path:
        return self.step06_dir / "city_morphology_metrics.parquet"

    @property
    def local_morphology_metrics(self) -> Path:
        return self.step06_dir / "local_morphology_metrics.parquet"

    @property
    def city_quality(self) -> Path:
        return self.step07_dir / "city_quality_scores.csv"

    @property
    def quality_review(self) -> Path:
        return self.step07_dir / "quality_review_checklist.csv"

    @property
    def city_profile_base(self) -> Path:
        return self.step08_dir / "city_profile_base.parquet"

    @property
    def scale_signature_long(self) -> Path:
        return self.step08_dir / "scale_signatures_long.parquet"

    @property
    def type_centers(self) -> Path:
        return self.step09_dir / "type_centers.csv"

    @property
    def type_contrib(self) -> Path:
        return self.step09_dir / "cluster_feature_contributions.csv"

    @property
    def city_type_profiles(self) -> Path:
        return self.step09_dir / "city_type_profiles.csv"

    @property
    def local_morphotypes(self) -> Path:
        return self.step09_dir / "local_morphotypes.parquet"

    @property
    def od_metrics(self) -> Path:
        return self.step10_dir / "OD_detour_metrics.parquet"

    @property
    def od_samples(self) -> Path:
        return self.step10_dir / "OD_detour_samples.parquet"

    @property
    def accessibility_metrics(self) -> Path:
        return self.step10_dir / "facility_accessibility_metrics.parquet"

    @property
    def inequality_metrics(self) -> Path:
        return self.step10_dir / "accessibility_inequality_metrics.parquet"

    @property
    def facility_classification(self) -> Path:
        return self.step10_dir / "facility_classification.parquet"

    @property
    def low_confidence_sensitivity(self) -> Path:
        return self.step10_dir / "low_confidence_sensitivity.csv"

    @property
    def step10_qc(self) -> Path:
        return self.step10_dir / "step10_quality_checks.csv"

    @property
    def model_a_table(self) -> Path:
        return self.step11_dir / "model_table_A_OD.csv"

    @property
    def model_b_table(self) -> Path:
        return self.step11_dir / "model_table_B_facility_accessibility.csv"

    @property
    def model_c_table(self) -> Path:
        return self.step11_dir / "model_table_C_inequality.csv"

    @property
    def model_coefficients(self) -> Path:
        return self.step11_dir / "model_coefficients.csv"

    @property
    def model_results(self) -> Path:
        return self.step11_dir / "table_model_results.csv"

    @property
    def robustness_results(self) -> Path:
        return self.step11_dir / "table_robustness_results.csv"

    @property
    def robustness_summary(self) -> Path:
        return self.step11_dir / "robustness_summary.csv"

    @property
    def marginal_effects(self) -> Path:
        return self.step11_dir / "model_marginal_effects.csv"

    @property
    def fig7_source(self) -> Path:
        return self.step11_dir / "figure_data_accessibility_by_type.csv"

    @property
    def fig8_source(self) -> Path:
        return self.step11_dir / "figure_data_morphotype_circuity.csv"

    @property
    def city_inequality_source(self) -> Path:
        return self.step11_dir / "figure_data_inequality.csv"

    @property
    def step11_qc(self) -> Path:
        return self.step11_dir / "step11_quality_checks.csv"

    @property
    def step12_qc(self) -> Path:
        return self.step12_dir / "step12_quality_checks.csv"

    @property
    def step12_readme(self) -> Path:
        return self.step12_dir / "README.md"

    @property
    def step12_run_log(self) -> Path:
        return self.step12_dir / "step12_run_log.json"

    @property
    def step12_repro(self) -> Path:
        return self.step12_dir / "reproducibility_status_manifest.json"

    @property
    def step12_figure_manifest(self) -> Path:
        return self.step12_dir / "step12_figure_manifest.csv"

    @property
    def step12_lineage(self) -> Path:
        return self.step12_dir / "step12_figure_data_lineage.csv"

    @property
    def step12_fig1_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig1_global_sample_cities_quality_tiers.csv"

    @property
    def step12_fig2_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig2_osm_quality_confidence.csv"

    @property
    def step12_fig3_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig3_multiscale_network_signatures.csv"

    @property
    def step12_fig4_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig4_morphotype_atlas_signatures.csv"

    @property
    def step12_fig5_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig5_morphotype_radar_signatures.csv"

    @property
    def step12_fig6_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig6_scale_network_transition_matrix.csv"

    @property
    def step12_fig7_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig7_core_city_morphotype_stability_and_performance.csv"

    @property
    def step12_fig8_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig8_accessibility_detour_performance_by_morphotype.csv"

    @property
    def step12_fig9_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig9_observed_weighted_group_means.csv"

    @property
    def step12_fig10_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig10_city_level_accessibility_inequality.csv"

    @property
    def step12_fig11_data(self) -> Path:
        return self.step12_dir / "figure_data_Fig11_model_coefficient_forest.csv"

    @property
    def manuscript(self) -> Path:
        return self.docs_dir / "manuscript.docx"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root",
        type=Path,
        default=Path(__file__).resolve().parents[2],
        help="Project root directory.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help=f"Output directory. Defaults to data/{STEP_NAME}; custom paths must stay inside that directory.",
    )
    parser.add_argument("--overwrite", action="store_true", help="Allow overwriting Step13 outputs.")
    return parser.parse_args()


def is_relative_to(path: Path, base: Path) -> bool:
    try:
        path.relative_to(base)
        return True
    except ValueError:
        return False


def resolve_step13_output_dir(root: Path, output_dir: Path | None) -> Path:
    allowed_root = (root / "data" / STEP_NAME).resolve()
    if output_dir is None:
        return allowed_root
    candidate = output_dir.expanduser()
    if not candidate.is_absolute():
        candidate = root / candidate
    out = candidate.resolve()
    if out != allowed_root and not is_relative_to(out, allowed_root):
        raise ValueError(
            f"Unsafe --output-dir: {out}. Step13 outputs may only be written to "
            f"{allowed_root} or one of its subdirectories."
        )
    return out


def get_paths(root: Path, output_dir: Path | None) -> StepPaths:
    root = root.resolve()
    out = resolve_step13_output_dir(root, output_dir)
    return StepPaths(
        root=root,
        output_dir=out,
        step01_dir=root / "data" / STEP01_NAME,
        step06_dir=root / "data" / STEP06_NAME,
        step07_dir=root / "data" / STEP07_NAME,
        step08_dir=root / "data" / STEP08_NAME,
        step09_dir=root / "data" / STEP09_NAME,
        step10_dir=root / "data" / STEP10_NAME,
        step11_dir=root / "data" / STEP11_NAME,
        step12_dir=root / "data" / STEP12_NAME,
        docs_dir=root / "docs",
    )


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def json_ready(obj: Any) -> Any:
    if obj is pd.NA:
        return None
    if isinstance(obj, Path):
        return str(obj)
    if isinstance(obj, (pd.Timestamp, datetime)):
        return obj.isoformat()
    if isinstance(obj, np.generic):
        return obj.item()
    if isinstance(obj, float) and (math.isnan(obj) or math.isinf(obj)):
        return None
    if isinstance(obj, dict):
        return {str(k): json_ready(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple, set)):
        return [json_ready(v) for v in obj]
    return obj


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(json_ready(obj), ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8-sig")


def safe_rel(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root.resolve()).as_posix()
    except Exception:
        return str(path)


def sha256_file(path: Path) -> str | None:
    if not path.exists() or not path.is_file():
        return None
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024 * 8), b""):
            h.update(chunk)
    return h.hexdigest()


def csv_shape(path: Path) -> tuple[int | None, int | None]:
    try:
        header = pd.read_csv(path, nrows=0)
        cols = len(header.columns)
        rows = 0
        for chunk in pd.read_csv(path, chunksize=100_000):
            rows += len(chunk)
        return rows, cols
    except Exception:
        pass
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if header is None:
                return 0, 0
            return sum(1 for _ in reader), len(header)
    except Exception:
        pass
    try:
        cols = len(pd.read_csv(path, nrows=0).columns)
    except Exception:
        cols = None
    try:
        with path.open("rb") as f:
            rows = max(sum(1 for _ in f) - 1, 0)
        return rows, cols
    except Exception:
        return None, cols


def parquet_shape(path: Path) -> tuple[int | None, int | None]:
    try:
        pf = pq.ParquetFile(path)
        return int(pf.metadata.num_rows), len(pf.schema_arrow.names)
    except Exception:
        return None, None


def xlsx_shape(path: Path) -> tuple[int | None, int | None, int | None, str]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        rows_total = 0
        max_cols = 0
        parts: list[str] = []
        for ws in wb.worksheets:
            rows_total += ws.max_row or 0
            max_cols = max(max_cols, ws.max_column or 0)
            parts.append(f"{ws.title}:{ws.max_row}x{ws.max_column}")
        sheet_count = len(wb.worksheets)
        wb.close()
        return rows_total, max_cols, sheet_count, "; ".join(parts)
    except Exception:
        return None, None, None, ""


def xlsx_sheet_names(path: Path) -> list[str]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        names = list(wb.sheetnames)
        wb.close()
        return names
    except Exception:
        return []


def text_shape(path: Path) -> tuple[int | None, int | None]:
    try:
        with path.open("r", encoding="utf-8") as f:
            return sum(1 for _ in f), 1
    except Exception:
        return None, None


def file_metadata(path: Path, root: Path, *, artifact_role: str = "", source_step: str = "", source_file: str = "") -> dict[str, Any]:
    path = path.resolve()
    exists = path.exists()
    rows: int | None = None
    columns: int | None = None
    sheet_count: int | None = None
    sheet_summary = ""
    if exists and path.is_file():
        suffix = path.suffix.lower()
        if suffix == ".csv":
            rows, columns = csv_shape(path)
        elif suffix == ".parquet":
            rows, columns = parquet_shape(path)
        elif suffix == ".xlsx":
            rows, columns, sheet_count, sheet_summary = xlsx_shape(path)
        elif suffix in {".md", ".json", ".txt"}:
            rows, columns = text_shape(path)
    stat = path.stat() if exists else None
    return {
        "artifact_name": path.name,
        "artifact_role": artifact_role,
        "relative_path": safe_rel(path, root),
        "absolute_path": str(path),
        "exists": exists,
        "rows": rows,
        "columns": columns,
        "sheet_count": sheet_count,
        "sheet_summary": sheet_summary,
        "size_bytes": stat.st_size if stat else None,
        "mtime": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds") if stat else None,
        "hash_sha256": sha256_file(path) if exists and path.is_file() else None,
        "source_step": source_step,
        "source_file": source_file,
    }


def collect_dir_manifest(root: Path, dirs: list[Path]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for base in dirs:
        if not base.exists():
            records.append({"relative_path": safe_rel(base, root), "exists": False})
            continue
        for p in sorted(x for x in base.rglob("*") if x.is_file()):
            stat = p.stat()
            records.append(
                {
                    "relative_path": safe_rel(p, root),
                    "size_bytes": stat.st_size,
                    "mtime_ns": stat.st_mtime_ns,
                }
            )
    return records


def manifest_digest(records: list[dict[str, Any]]) -> str:
    payload = json.dumps(records, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def compare_manifests(before: list[dict[str, Any]], after: list[dict[str, Any]]) -> list[str]:
    before_map = {r.get("relative_path"): r for r in before}
    after_map = {r.get("relative_path"): r for r in after}
    changed: list[str] = []
    for key in sorted(set(before_map) | set(after_map)):
        if before_map.get(key) != after_map.get(key):
            changed.append(str(key))
    return changed


def check_existing_outputs(paths: StepPaths, overwrite: bool) -> None:
    if overwrite:
        return
    existing = [name for name in OUTPUT_FILES if (paths.output_dir / name).exists()]
    if existing:
        preview = ", ".join(existing[:8])
        raise FileExistsError(f"Step13 outputs already exist ({preview}). Re-run with --overwrite.")


def read_csv(path: Path, **kwargs: Any) -> pd.DataFrame:
    return pd.read_csv(path, **kwargs)


def read_parquet(path: Path, columns: list[str] | None = None) -> pd.DataFrame:
    return pd.read_parquet(path, columns=columns)


def existing_parquet_columns(path: Path, columns: list[str]) -> list[str]:
    schema_names = set(pq.ParquetFile(path).schema_arrow.names)
    return [c for c in columns if c in schema_names]


def read_model_b_city_ids(path: Path) -> set[str]:
    cols = ["city_id", "main_analysis_included", "analysis_scope", "scale"]
    df = pd.read_csv(path, usecols=lambda c: c in cols)
    mask = pd.Series(True, index=df.index)
    if "main_analysis_included" in df.columns:
        mask &= df["main_analysis_included"].fillna(False).astype(bool)
    if "analysis_scope" in df.columns:
        mask &= df["analysis_scope"].eq("main_fit")
    if "scale" in df.columns:
        mask &= df["scale"].eq("hex_1km")
    return set(df.loc[mask, "city_id"].dropna().astype(str).unique())


def attach_english_labels(df: pd.DataFrame, key_col: str = "morphotype") -> pd.DataFrame:
    out = df.copy()
    if key_col in out.columns:
        out["morphotype_label_en"] = out[key_col].map(MORPHOTYPE_LABELS_EN)
        out["morphotype_key_policy"] = "MT01-MT10 are the unique keys; descriptive names may repeat."
    return out


def fig5_axis_scores(fig5: pd.DataFrame, type_centers: pd.DataFrame) -> pd.DataFrame:
    required = {"morphotype", "axis_code", "score_metric_label", "score_value"}
    if required.issubset(fig5.columns):
        axis = fig5[["morphotype", "axis_code", "score_metric_label", "score_value"]].copy()
        axis["score_value"] = pd.to_numeric(axis["score_value"], errors="coerce")
        return axis

    # Fallback follows the current Fig5 axis order: D is Hierarchy and E is Cul-de-sac.
    score_specs = [
        ("A", "density_score", "Density"),
        ("B", "fine_grain_score", "Fine grain"),
        ("C", "grid_score", "Grid"),
        ("D", "hierarchy_score", "Hierarchy"),
        ("E", "culdesac_score", "Cul-de-sac"),
        ("F", "circuity_score", "Circuity"),
        ("G", "segment_length_score", "Segment length"),
    ]
    rows: list[dict[str, Any]] = []
    for _, r in type_centers.iterrows():
        mt = str(r.get("morphotype", ""))
        for axis_code, metric, label in score_specs:
            rows.append(
                {
                    "morphotype": mt,
                    "axis_code": axis_code,
                    "score_metric_label": label,
                    "score_value": pd.to_numeric(pd.Series([r.get(metric)]), errors="coerce").iloc[0],
                }
            )
    return pd.DataFrame(rows)


def format_score_axis(row: pd.Series) -> str:
    value = row.get("score_value")
    value_text = f"{float(value):+.2f}" if pd.notna(value) else "NA"
    return f"{row.get('axis_code')} {row.get('score_metric_label')} ({value_text})"


def structural_signature_from_scores(axis_rows: pd.DataFrame) -> str:
    rows = axis_rows.copy()
    rows["score_value"] = pd.to_numeric(rows["score_value"], errors="coerce")
    rows = rows.dropna(subset=["score_value"])
    if rows.empty:
        return "Local morphotype; structural score pattern is unavailable."

    high = rows[rows["score_value"].ge(0.45)].sort_values("score_value", ascending=False).head(3)
    low = rows[rows["score_value"].le(-0.45)].sort_values("score_value", ascending=True).head(3)
    parts: list[str] = []
    if len(high):
        parts.append("higher " + ", ".join(high.apply(format_score_axis, axis=1)))
    if len(low):
        parts.append("lower " + ", ".join(low.apply(format_score_axis, axis=1)))
    if not parts:
        leading = rows.sort_values("score_value", ascending=False).head(3)
        trailing = rows.sort_values("score_value", ascending=True).head(2)
        parts.append("mixed/near-neutral profile led by " + ", ".join(leading.apply(format_score_axis, axis=1)))
        parts.append("comparatively lower " + ", ".join(trailing.apply(format_score_axis, axis=1)))
    return "Local units with " + "; ".join(parts) + "."


def build_morphotype_totals(local_morphotypes: pd.DataFrame) -> pd.DataFrame:
    df = local_morphotypes.copy()
    if "network_type" in df.columns:
        df = df[df["network_type"].astype(str).eq("drive")]
    if "scale" in df.columns:
        df = df[df["scale"].astype(str).eq("hex_1km")]

    population_col = ""
    for candidate in ["population_sum", "population_weight"]:
        if candidate in df.columns and pd.to_numeric(df[candidate], errors="coerce").fillna(0).sum() > 0:
            population_col = candidate
            break

    if population_col:
        df["_population_basis"] = pd.to_numeric(df[population_col], errors="coerce").fillna(0)
    else:
        df["_population_basis"] = 1.0
        population_col = "unit_count_fallback"

    if "type_confidence" in df.columns:
        df["_type_confidence"] = pd.to_numeric(df["type_confidence"], errors="coerce")
    else:
        df["_type_confidence"] = np.nan

    total_units = len(df)
    total_population = float(df["_population_basis"].sum())
    grouped = (
        df.groupby("morphotype", dropna=False)
        .agg(
            unit_count_total=("morphotype", "size"),
            population_total=("_population_basis", "sum"),
            mean_type_confidence=("_type_confidence", "mean"),
        )
        .reset_index()
    )
    grouped["unit_share_total"] = grouped["unit_count_total"] / total_units if total_units else np.nan
    grouped["population_share_total"] = (
        grouped["population_total"] / total_population if total_population > 0 else np.nan
    )
    grouped["total_share_basis"] = (
        f"Step09 local units, network_type=drive, scale=hex_1km; population basis={population_col}."
    )
    return grouped


def morphotype_name_lookup(type_centers: pd.DataFrame | None) -> dict[str, str]:
    if type_centers is None or not {"morphotype", "morphotype_name"}.issubset(type_centers.columns):
        return {mt: mt for mt in MORPHOTYPES}
    names = (
        type_centers[["morphotype", "morphotype_name"]]
        .dropna(subset=["morphotype"])
        .drop_duplicates("morphotype")
        .set_index("morphotype")["morphotype_name"]
        .astype(str)
        .to_dict()
    )
    return {mt: names.get(mt, mt) for mt in MORPHOTYPES}


def morphotype_share_columns(df: pd.DataFrame, metric: str) -> dict[str, str]:
    candidates = [
        {mt: f"{metric}_share__{mt}" for mt in MORPHOTYPES},
        {mt: f"{metric}_share_{mt}" for mt in MORPHOTYPES},
    ]
    columns = set(df.columns)
    for mapping in candidates:
        if set(mapping.values()).issubset(columns):
            return mapping
    partial = {mt: col for mt in MORPHOTYPES for col in [f"{metric}_share__{mt}", f"{metric}_share_{mt}"] if col in columns}
    return partial if len(partial) == len(MORPHOTYPES) else {}


def morphotype_share_matrix(df: pd.DataFrame, metric: str) -> pd.DataFrame:
    mapping = morphotype_share_columns(df, metric)
    if not mapping:
        return pd.DataFrame(index=df.index)
    return pd.DataFrame(
        {mt: pd.to_numeric(df[col], errors="coerce") for mt, col in mapping.items()},
        index=df.index,
    )


def dominant_from_share_matrix(matrix: pd.DataFrame) -> tuple[pd.Series, pd.Series]:
    if matrix.empty:
        return pd.Series(pd.NA, index=matrix.index, dtype="object"), pd.Series(np.nan, index=matrix.index)
    has_any = matrix.notna().any(axis=1)
    dominant = matrix.idxmax(axis=1).astype("object").where(has_any, pd.NA)
    share = matrix.max(axis=1, skipna=True).where(has_any, np.nan)
    return dominant, share


def lookup_morphotype_share(matrix: pd.DataFrame, morphotypes: pd.Series) -> pd.Series:
    if matrix.empty:
        return pd.Series(np.nan, index=morphotypes.index)
    col_pos = {mt: i for i, mt in enumerate(matrix.columns)}
    values = matrix.to_numpy(dtype=float, copy=False)
    out: list[float] = []
    for row_pos, mt in enumerate(morphotypes.astype("object").tolist()):
        pos = col_pos.get(str(mt))
        out.append(float(values[row_pos, pos]) if pos is not None else np.nan)
    return pd.Series(out, index=morphotypes.index)


def add_dominant_morphotype_fields(df: pd.DataFrame, type_centers: pd.DataFrame | None) -> pd.DataFrame:
    out = df.copy()
    names = morphotype_name_lookup(type_centers)
    population_matrix = morphotype_share_matrix(out, "population")
    unit_matrix = morphotype_share_matrix(out, "unit")
    area_matrix = morphotype_share_matrix(out, "area")

    population_dominant, population_share = dominant_from_share_matrix(population_matrix)
    unit_dominant, unit_share = dominant_from_share_matrix(unit_matrix)

    out["population_dominant_morphotype"] = population_dominant
    out["population_dominant_morphotype_name"] = population_dominant.map(names)
    out["population_dominant_share"] = population_share
    out["unit_dominant_morphotype"] = unit_dominant
    out["unit_dominant_morphotype_name"] = unit_dominant.map(names)
    out["unit_dominant_share"] = unit_share

    # Backward-compatible fields are now explicitly population-dominant.
    out["dominant_morphotype"] = out["population_dominant_morphotype"]
    out["dominant_morphotype_name"] = out["population_dominant_morphotype_name"]
    out["dominant_population_share"] = out["population_dominant_share"]
    out["dominant_unit_share"] = lookup_morphotype_share(unit_matrix, population_dominant)
    if not area_matrix.empty and "dominant_area_share" in out.columns:
        out["dominant_area_share"] = lookup_morphotype_share(area_matrix, population_dominant)
    return out


def build_appendix_city_type_profiles(city_type_profiles: pd.DataFrame, type_centers: pd.DataFrame) -> pd.DataFrame:
    out = add_dominant_morphotype_fields(city_type_profiles, type_centers)
    out["dominant_field_semantics_note"] = DOMINANT_COMPATIBILITY_NOTE
    leading = [
        "city_id",
        "city_name_en",
        "country",
        "iso3",
        "region",
        "sample_group",
        "quality_tier",
        "network_type",
        "scale",
        "profile_scope",
        "population_dominant_morphotype",
        "population_dominant_morphotype_name",
        "population_dominant_share",
        "unit_dominant_morphotype",
        "unit_dominant_morphotype_name",
        "unit_dominant_share",
        "dominant_morphotype",
        "dominant_morphotype_name",
        "dominant_population_share",
        "dominant_unit_share",
        "dominant_area_share",
    ]
    ordered = [c for c in leading if c in out.columns]
    ordered += [c for c in out.columns if c not in ordered]
    return out[ordered]


def build_city_morphotype_composition(city_type_profiles: pd.DataFrame, type_centers: pd.DataFrame) -> pd.DataFrame:
    df = city_type_profiles.copy()
    df = df[df["network_type"].astype(str).eq("drive") & df["scale"].astype(str).eq("hex_1km")].copy()
    base_cols = [
        "city_id",
        "city_name_en",
        "country",
        "iso3",
        "region",
        "sample_group",
        "quality_tier",
        "network_type",
        "scale",
        "profile_scope",
        "morphotype_entropy",
        "mean_type_confidence",
        "n_units",
    ]
    out = pd.DataFrame(index=df.index)
    for col in base_cols:
        out[col] = df[col] if col in df.columns else np.nan
    for mt in MORPHOTYPES:
        out[f"population_share_{mt}"] = pd.to_numeric(df.get(f"population_share__{mt}"), errors="coerce")
    for mt in MORPHOTYPES:
        out[f"unit_share_{mt}"] = pd.to_numeric(df.get(f"unit_share__{mt}"), errors="coerce")
    out = add_dominant_morphotype_fields(out, type_centers)
    out["composition_scope_note"] = (
        "Rows use Step09 city profiles with network_type=drive and scale=hex_1km; "
        "the upstream profile table contains non_excluded_candidates records for available cities. "
        + DOMINANT_COMPATIBILITY_NOTE
    )
    sort_cols = [c for c in ["sample_group", "quality_tier", "region", "city_id"] if c in out.columns]
    leading = [
        "city_id",
        "city_name_en",
        "country",
        "iso3",
        "region",
        "sample_group",
        "quality_tier",
        "network_type",
        "scale",
        "profile_scope",
        "population_dominant_morphotype",
        "population_dominant_morphotype_name",
        "population_dominant_share",
        "unit_dominant_morphotype",
        "unit_dominant_morphotype_name",
        "unit_dominant_share",
        "dominant_morphotype",
        "dominant_morphotype_name",
        "dominant_population_share",
        "dominant_unit_share",
        "morphotype_entropy",
        "mean_type_confidence",
        "n_units",
    ]
    share_cols = [f"population_share_{mt}" for mt in MORPHOTYPES] + [f"unit_share_{mt}" for mt in MORPHOTYPES]
    ordered = [c for c in leading + share_cols + ["composition_scope_note"] if c in out.columns]
    ordered += [c for c in out.columns if c not in ordered]
    return out[ordered].sort_values(sort_cols).reset_index(drop=True)


def build_morphotype_definitions(
    type_centers: pd.DataFrame,
    fig5: pd.DataFrame,
    city_composition: pd.DataFrame,
    local_morphotypes: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    axis = fig5_axis_scores(fig5, type_centers)
    totals = build_morphotype_totals(local_morphotypes)
    core = city_composition[
        city_composition["quality_tier"].astype(str).eq("core")
        & city_composition["network_type"].astype(str).eq("drive")
        & city_composition["scale"].astype(str).eq("hex_1km")
    ].copy()

    center_names = (
        type_centers[["morphotype", "morphotype_name"]].drop_duplicates("morphotype")
        if {"morphotype", "morphotype_name"}.issubset(type_centers.columns)
        else pd.DataFrame({"morphotype": MORPHOTYPES, "morphotype_name": [np.nan] * len(MORPHOTYPES)})
    )
    center_names = center_names.set_index("morphotype")["morphotype_name"].to_dict()
    totals = totals.set_index("morphotype")

    rows: list[dict[str, Any]] = []
    rep_rows: list[dict[str, Any]] = []
    for mt in MORPHOTYPES:
        mt_axis = axis[axis["morphotype"].astype(str).eq(mt)].copy()
        mt_axis["score_value"] = pd.to_numeric(mt_axis["score_value"], errors="coerce")
        dominant_axes = "; ".join(
            mt_axis.sort_values("score_value", ascending=False).head(3).apply(format_score_axis, axis=1)
        )
        pop_col = f"population_share_{mt}"
        unit_col = f"unit_share_{mt}"
        ranked = core.copy()
        ranked["mt_population_share"] = pd.to_numeric(ranked.get(pop_col), errors="coerce").fillna(0)
        ranked["mt_unit_share"] = pd.to_numeric(ranked.get(unit_col), errors="coerce").fillna(0)
        ranked = ranked.sort_values(
            ["mt_population_share", "city_name_en", "city_id"],
            ascending=[False, True, True],
        ).head(5)
        rep_text = "; ".join(
            f"{r.city_name_en} ({r.mt_population_share:.3f})" for r in ranked.itertuples(index=False)
        )
        dominant_col = "population_dominant_morphotype" if "population_dominant_morphotype" in core.columns else "dominant_morphotype"
        n_dominant = int(core[dominant_col].astype(str).eq(mt).sum())
        note = "All representatives are selected from core drive hex_1km profiles."
        if n_dominant == 0:
            note = (
                f"No core drive hex_1km city has {mt} as population_dominant_morphotype; "
                "representatives are the highest population-share core cities for this local morphotype."
            )
        elif n_dominant < 5:
            note = (
                f"Only {n_dominant} core drive hex_1km city profile(s) have {mt} as population_dominant_morphotype; "
                "the top-five list also includes high-share non-dominant cities."
            )
        if len(ranked) < 5:
            note += f" Only {len(ranked)} core profile(s) were available."

        total_row = totals.loc[mt] if mt in totals.index else pd.Series(dtype=object)
        rows.append(
            {
                "morphotype": mt,
                "morphotype_name_existing": center_names.get(mt),
                "english_label": MORPHOTYPE_LABELS_EN.get(mt),
                "structural_signature": structural_signature_from_scores(mt_axis),
                "dominant_axes": dominant_axes,
                "representative_core_cities": rep_text,
                "n_dominant_core_cities": n_dominant,
                "population_share_total": total_row.get("population_share_total", np.nan),
                "unit_share_total": total_row.get("unit_share_total", np.nan),
                "mean_type_confidence": total_row.get("mean_type_confidence", np.nan),
                "selection_rule": (
                    "Representatives: top five quality_tier=core, network_type=drive, scale=hex_1km "
                    f"city profiles ranked by population_share_{mt}. Totals: "
                    f"{total_row.get('total_share_basis', 'Step09 local units, drive hex_1km.')}"
                ),
                "note": note,
            }
        )
        for rank, r in enumerate(ranked.itertuples(index=False), start=1):
            rep_rows.append(
                {
                    "morphotype": mt,
                    "rank": rank,
                    "city_id": getattr(r, "city_id", ""),
                    "city_name_en": getattr(r, "city_name_en", ""),
                    "country": getattr(r, "country", ""),
                    "iso3": getattr(r, "iso3", ""),
                    "region": getattr(r, "region", ""),
                    "population_share": getattr(r, "mt_population_share", np.nan),
                    "unit_share": getattr(r, "mt_unit_share", np.nan),
                    "city_dominant_morphotype": getattr(r, "dominant_morphotype", ""),
                    "city_population_dominant_morphotype": getattr(r, "population_dominant_morphotype", ""),
                    "city_unit_dominant_morphotype": getattr(r, "unit_dominant_morphotype", ""),
                    "is_city_dominant_for_type": str(getattr(r, "dominant_morphotype", "")) == mt,
                    "is_city_population_dominant_for_type": str(getattr(r, "population_dominant_morphotype", "")) == mt,
                    "is_city_unit_dominant_for_type": str(getattr(r, "unit_dominant_morphotype", "")) == mt,
                    "quality_tier": getattr(r, "quality_tier", ""),
                    "network_type": getattr(r, "network_type", ""),
                    "scale": getattr(r, "scale", ""),
                }
            )

    definitions = pd.DataFrame(rows)
    representatives = pd.DataFrame(rep_rows)
    return definitions, representatives


def build_table2(city_master: pd.DataFrame, core_city_ids: set[str]) -> pd.DataFrame:
    df = city_master.copy()
    df["core_analysis_flag"] = df["city_id"].astype(str).isin(core_city_ids)
    rows = []
    for region, g in df.groupby("region", dropna=False):
        total = len(g)
        rows.append(
            {
                "region": region,
                "main_sample_city_count": int(g["sample_group"].eq("main_80").sum()),
                "china_pressure_test_city_count": int(g["sample_group"].eq("china_pressure_test").sum()),
                "total_city_count": total,
                "core_analysis_city_count": int(g["core_analysis_flag"].sum()),
                "share_of_all_cities": total / len(df) if len(df) else np.nan,
                "city_examples": "; ".join(g["city_name_en"].astype(str).head(5)),
                "source_step": STEP01_NAME,
                "source_file": "data/01_city_boundaries_and_sample_inventory/city_sample/city_master.parquet",
            }
        )
    rows.append(
        {
            "region": "ALL",
            "main_sample_city_count": int(df["sample_group"].eq("main_80").sum()),
            "china_pressure_test_city_count": int(df["sample_group"].eq("china_pressure_test").sum()),
            "total_city_count": int(len(df)),
            "core_analysis_city_count": int(df["core_analysis_flag"].sum()),
            "share_of_all_cities": 1.0,
            "city_examples": "",
            "source_step": STEP01_NAME,
            "source_file": "data/01_city_boundaries_and_sample_inventory/city_sample/city_master.parquet",
        }
    )
    return pd.DataFrame(rows).sort_values(["region"]).reset_index(drop=True)


def build_table3(city_master: pd.DataFrame, core_city_ids: set[str]) -> pd.DataFrame:
    df = city_master.copy()
    df["core_analysis_flag"] = df["city_id"].astype(str).isin(core_city_ids)
    group_cols = ["sample_group", "region", "morphology_prior"]
    out = (
        df.groupby(group_cols, dropna=False)
        .agg(
            city_count=("city_id", "nunique"),
            core_analysis_city_count=("core_analysis_flag", "sum"),
            city_examples=("city_name_en", lambda s: "; ".join(s.astype(str).head(4))),
        )
        .reset_index()
    )
    out.insert(0, "stratum_id", [f"S{i:03d}" for i in range(1, len(out) + 1)])
    out["selection_logic"] = (
        "Cities are fixed by the Step01 sample master; Step13 summarizes strata and does not resample."
    )
    out["source_step"] = STEP01_NAME
    out["source_file"] = "data/01_city_boundaries_and_sample_inventory/city_sample/city_master.parquet"
    return out


def build_data_source_records(paths: StepPaths, input_meta: dict[str, dict[str, Any]]) -> pd.DataFrame:
    source_specs = [
        ("city_boundary_sample", "GHS-UCDB R2024A / UCDB 2025", "Urban boundary and city sample anchor", paths.city_master),
        ("road_network", "OSM fixed extracts / cleaned drive and walk networks", "Morphology metrics and route validation", paths.city_morphology_metrics),
        ("population", "GHSL population grids, supplemented by WorldPop summaries", "Population weights and city controls", paths.city_profile_base),
        ("built_environment", "GHSL built-up / WSF support indicators", "Built support and quality controls", paths.city_profile_base),
        ("facility_poi", "OSM facilities plus Overture Places", "15-minute accessibility facility categories", paths.facility_classification),
        ("osm_quality_history", "ohsome and OSM history summaries", "OSM quality scoring and confidence weights", paths.city_quality),
        ("morphotype_assignment", "Step09 KMeans morphotype outputs", "MT01-MT10 type centers and unit assignments", paths.type_centers),
        ("accessibility_detour", "Step10 accessibility and OD detour outputs", "Validation outcomes and appendix E data", paths.od_metrics),
        ("model_results", "Step11 model and robustness outputs", "Validation model coefficients and robustness appendix", paths.model_coefficients),
        ("figure_products", "Step12 publication figure outputs", "Figure data and figure file provenance", paths.step12_figure_manifest),
    ]
    rows = []
    for data_domain, source_name, role, path in source_specs:
        rel = safe_rel(path, paths.root)
        meta = input_meta.get(rel, {})
        rows.append(
            {
                "data_domain": data_domain,
                "source_name": source_name,
                "project_role": role,
                "representative_file": rel,
                "rows": meta.get("rows"),
                "columns": meta.get("columns"),
                "hash_sha256": meta.get("hash_sha256"),
                "mtime": meta.get("mtime"),
                "source_step": meta.get("source_step", ""),
                "version_or_snapshot_note": "See upstream step README/reproducibility JSON for exact extraction timestamps and software state.",
                "excel_export_policy": "Full table in appendix only when aggregated or moderate-size; very large tables are material-list only.",
            }
        )
    return pd.DataFrame(rows)


def build_table5(paths: StepPaths) -> pd.DataFrame:
    steps = [
        ("00", "environment_manifest", "Record environment versions and download task status", "data/00_environment_versions_and_data_manifest"),
        ("01", "city_sample", "Fix city boundaries and 86-city sample master", "data/01_city_boundaries_and_sample_inventory"),
        ("02", "road_network_download", "Download OSM road network extracts", "data/02_download_road_network_data"),
        ("03", "population_built_download", "Download population and built-environment rasters", "data/03_download_population_built_environment_data"),
        ("04", "facility_quality_download", "Download OSM/Overture facilities and quality history", "data/04_download_facility_and_quality_history_data"),
        ("05", "spatial_units", "Build full-city, core, and grid spatial units", "data/05_build_multiscale_spatial_units"),
        ("06", "morphology_metrics", "Clean networks and compute morphology metrics", "data/06_clean_networks_compute_morphology_metrics"),
        ("07", "quality_confidence", "Generate city quality scores and confidence weights", "data/07_generate_quality_scores_and_type_confidence"),
        ("08", "scale_signatures", "Generate multiscale signatures and city profiles", "data/08_generate_scale_signatures_and_city_profiles"),
        ("09", "morphotype_clustering", "Identify MT01-MT10 morphotypes", "data/09_cluster_morphotypes"),
        ("10", "access_detour_validation", "Compute accessibility and OD detour validation metrics", "data/10_calculate_accessibility_and_detour_validation_metrics"),
        ("11", "models_robustness", "Fit explanatory models and robustness checks", "data/11_statistical_models_and_robustness_checks"),
        ("12", "paper_figures", "Draw manuscript figures", "data/12_draw_paper_figures"),
        ("13", "tables_appendices", "Export manuscript tables, appendices, provenance, and QC", "data/13_export_paper_tables_appendices"),
    ]
    rows = []
    for step_id, module_id, purpose, rel_dir in steps:
        d = paths.root / rel_dir
        rows.append(
            {
                "step_id": step_id,
                "module_id": module_id,
                "purpose": purpose,
                "output_directory": rel_dir,
                "directory_exists": d.exists(),
                "file_count": sum(1 for p in d.rglob("*") if p.is_file()) if d.exists() else 0,
                "step13_policy": "read_only_upstream" if step_id != "13" else "write_step13_outputs_only",
            }
        )
    return pd.DataFrame(rows)


def build_table6() -> pd.DataFrame:
    rows = [
        ("sample_frame", "86 fixed cities", "city_id", "Step01 city master", "Table2, Table3, Appendix B"),
        ("quality_confidence", "OSM quality scoring", "city_id", "Step07 quality scores", "Appendix C and model weights"),
        ("morphology_metrics", "Multiscale network morphology", "city_id, scale, unit_id", "Step06/08 metrics", "Morphotype clustering and Fig4"),
        ("morphotype_clustering", "MT01-MT10 urban street morphotypes", "morphotype", "Step09 type centers and assignments", "Appendix D and Fig5/Fig6"),
        ("facility_accessibility", "15-minute facility accessibility", "unit_id, category", "Step10 facility accessibility metrics", "Appendix E and Model B/C"),
        ("od_detour", "OD circuity and route validation", "city_id, morphotype pair", "Step10 OD detour metrics", "Appendix E and Model A/Fig8"),
        ("explanatory_models", "Model A/B/C explanatory associations", "model_id, term", "Step11 model outputs", "Table7, Appendix F"),
        ("robustness_checks", "High-confidence, unweighted, and small-group robustness", "scenario, model_id", "Step11 robustness outputs", "Table8, Appendix F"),
        ("figure_provenance", "Figure files and source data", "figure_id", "Step12 manifest and lineage", "Material checklist"),
        ("export_package", "Tables, appendices, provenance, reproducibility state", "artifact file", "Step13 exports", "Submission package"),
    ]
    return pd.DataFrame(
        rows,
        columns=[
            "analysis_module",
            "module_scope",
            "primary_key",
            "main_input",
            "paper_or_appendix_use",
        ],
    )


def build_table7(model_results: pd.DataFrame) -> pd.DataFrame:
    df = model_results.copy()
    if "scenario" in df.columns:
        df = df[df["scenario"].astype(str).str.startswith("main")]
    rows = []
    family_labels = {
        "A_OD_circuity": "Model A: OD detour / route performance",
        "B_facility_accessibility": "Model B: facility accessibility",
        "C_accessibility_inequality": "Model C: accessibility inequality",
    }
    outcome_labels = {
        "accessibility_gini": "Accessibility Gini",
        "no_access_population_share": "No-access population share",
        "population_weighted_access_share": "Population-weighted access share",
        "accessibility_score": "Accessibility score",
        "access_15min_num": "15-minute facility count",
        "log_od_circuity_weighted_mean": "Log weighted OD circuity",
        "od_circuity_weighted_mean": "Weighted OD circuity mean",
        "od_circuity_weighted_p50": "Weighted OD circuity p50",
        "od_circuity_weighted_p90": "Weighted OD circuity p90",
        "route_found_weighted_share": "Route-found weighted share",
    }
    for (model_family, outcome), g in df.groupby(["model_family", "outcome"], dropna=False):
        g_ok = g[g.get("status", "ok").eq("ok")] if "status" in g.columns else g
        first = g_ok.iloc[0] if len(g_ok) else g.iloc[0]
        term_kind = g["term_kind"] if "term_kind" in g.columns else pd.Series("", index=g.index)
        term = g["term"] if "term" in g.columns else pd.Series("", index=g.index)
        control_terms = sorted(term[term_kind.eq("control")].dropna().astype(str).unique())
        category_terms = sorted(term[term_kind.eq("facility_category")].dropna().astype(str).unique())
        has_interactions = bool(g["term"].dropna().astype(str).str.contains(":", regex=False).any()) if "term" in g.columns else False
        rows.append(
            {
                "model_block": family_labels.get(str(model_family), str(model_family)),
                "model_family": model_family,
                "outcome": outcome,
                "outcome_label": outcome_labels.get(str(outcome), str(outcome)),
                "model_ids": "; ".join(sorted(g["model_id"].dropna().astype(str).unique())),
                "coefficient_rows": int(len(g)),
                "morphotype_term_rows": int(g["term_kind"].eq("morphotype").sum()) if "term_kind" in g.columns else None,
                "facility_category_term_rows": int(g["term_kind"].eq("facility_category").sum()) if "term_kind" in g.columns else None,
                "control_term_rows": int(g["term_kind"].eq("control").sum()) if "term_kind" in g.columns else None,
                "control_variables": "; ".join(control_terms) if control_terms else "None or absorbed by fixed effects",
                "facility_category_terms": "; ".join(category_terms) if category_terms else "",
                "morphotype_category_interactions": "yes" if has_interactions else "no",
                "nobs": first.get("nobs"),
                "weight_col": first.get("weight_col"),
                "cluster_col": first.get("cluster_col"),
                "covariance_type": first.get("cov_type"),
                "formula": first.get("formula"),
                "model_c_inequality_outcome": "yes" if str(model_family) == "C_accessibility_inequality" else "no",
                "interpretation_scope": "Explanatory association; not a causal effect.",
                "source_step": STEP11_NAME,
                "source_file": "data/11_statistical_models_and_robustness_checks/table_model_results.csv",
            }
        )
    return pd.DataFrame(rows).sort_values(["model_family", "outcome"]).reset_index(drop=True)


def build_table8(robustness_summary: pd.DataFrame) -> pd.DataFrame:
    out = robustness_summary.copy()
    out["source_step"] = STEP11_NAME
    out["source_file"] = "data/11_statistical_models_and_robustness_checks/robustness_summary.csv"
    out["step11_small_group_warning_policy"] = "non_blocking_warn_recorded"
    out["interpretation_scope"] = "Robustness checks probe sensitivity of explanatory associations; they do not establish causality."
    return out


def variable_dictionary_rows() -> list[dict[str, Any]]:
    specs = [
        ("city_id", "identifier", "Stable project city identifier.", "Unique project-level city key.", "", "Step01"),
        ("city_name_en", "identifier", "English city name used in outputs.", "Display city name in English.", "", "Step01"),
        ("region", "sample", "World-region stratum used for sample summaries.", "Regional sample stratum.", "", "Step01"),
        ("sample_group", "sample", "Main 80-city sample or China pressure-test group.", "Sample membership flag.", "", "Step01"),
        ("morphology_prior", "sample", "Pre-analysis morphology prior used for sample balancing.", "Prior morphology label used during sample design.", "", "Step01"),
        ("quality_score", "quality", "Composite OSM quality score.", "Composite OSM quality score.", "0-100", "Step07"),
        ("quality_tier", "quality", "Quality tier used for core and sensitivity labels.", "Quality stratum.", "", "Step07"),
        ("quality_weight", "quality", "Weight derived from quality score for downstream summaries.", "Quality-derived downstream weight.", "", "Step07"),
        ("network_integrity_score", "quality", "Road-network integrity component.", "Road-network integrity subscore.", "0-100", "Step07"),
        ("historical_maturity_score", "quality", "OSM history maturity component.", "OSM history maturity subscore.", "0-100", "Step07"),
        ("poi_completeness_score", "quality", "Facility/POI completeness component.", "Facility POI completeness subscore.", "0-100", "Step07"),
        ("local_coverage_score", "quality", "Local metric coverage component.", "Local metric coverage subscore.", "0-100", "Step07"),
        ("morphotype", "morphotype", "Unique morphotype key MT01-MT10.", "Unique MT01-MT10 morphotype key.", "", "Step09"),
        ("morphotype_name", "morphotype", "Descriptive morphotype name; not guaranteed unique.", "Descriptive morphotype name carried from upstream outputs.", "", "Step09"),
        ("morphotype_label_en", "morphotype", "Unique English display label attached to MT key.", "Unique English label attached to the MT key.", "", "Step12/13"),
        ("english_label", "morphotype", "Paper-readable English label for a local morphotype.", "Paper-facing English label for the local morphotype.", "", "Step13"),
        ("structural_signature", "morphotype", "Concise structural description derived from Step09/Fig5 score axes.", "Structural signature derived from Step09/Fig5 score axes.", "", "Step13"),
        ("dominant_axes", "morphotype", "Top Fig5 A-G axes by Step09 score value; D=Hierarchy and E=Cul-de-sac.", "Top Fig5 A-G score axes ranked by Step09 score value.", "", "Step13"),
        ("representative_core_cities", "morphotype", "Top-five core drive hex_1km cities ranked by population share of this local morphotype.", "Top five core drive hex_1km city examples ranked by population share.", "", "Step13"),
        ("n_dominant_core_cities", "morphotype", "Number of core drive hex_1km city profiles whose population-dominant morphotype is the MT key.", "Count of core drive hex_1km city profiles where the MT key is population-dominant.", "count", "Step13"),
        ("type_confidence", "morphotype", "Local morphotype assignment confidence.", "Local type assignment confidence.", "0-1", "Step09"),
        ("cluster_probability", "morphotype", "Cluster assignment probability.", "Cluster membership probability.", "0-1", "Step09"),
        ("population_dominant_morphotype", "morphotype", "City profile's largest population-share local morphotype.", "Local morphotype with the highest population share in the city profile.", "", "Step13"),
        ("population_dominant_morphotype_name", "morphotype", "Display label for population_dominant_morphotype.", "Display label for the population-dominant morphotype.", "", "Step13"),
        ("population_dominant_share", "morphotype", "Population share of population_dominant_morphotype within a city profile.", "Population share of the population-dominant morphotype in the city profile.", "0-1", "Step13"),
        ("unit_dominant_morphotype", "morphotype", "City profile's largest spatial-unit-share local morphotype.", "Local morphotype with the highest spatial-unit share in the city profile.", "", "Step13"),
        ("unit_dominant_morphotype_name", "morphotype", "Display label for unit_dominant_morphotype.", "Display label for the unit-dominant morphotype.", "", "Step13"),
        ("unit_dominant_share", "morphotype", "Spatial-unit share of unit_dominant_morphotype within a city profile.", "Spatial-unit share of the unit-dominant morphotype in the city profile.", "0-1", "Step13"),
        ("dominant_morphotype", "morphotype", "Compatibility field equal to population_dominant_morphotype in Step13 outputs.", "Compatibility field equal to population_dominant_morphotype.", "", "Step13"),
        ("dominant_morphotype_name", "morphotype", "Compatibility field equal to population_dominant_morphotype_name in Step13 outputs.", "Compatibility field equal to population_dominant_morphotype_name.", "", "Step13"),
        ("dominant_population_share", "morphotype", "Compatibility field equal to population_dominant_share in Step13 outputs.", "Compatibility field equal to population_dominant_share.", "0-1", "Step13"),
        ("dominant_unit_share", "morphotype", "Spatial-unit share of dominant_morphotype, where dominant_morphotype is population-dominant in Step13 outputs; use unit_dominant_share for the unit-dominant share.", "Spatial-unit share for the population-dominant morphotype; use unit_dominant_share for unit dominance.", "0-1", "Step13"),
        ("morphotype_entropy", "morphotype", "Entropy of MT01-MT10 composition within a city profile.", "Entropy of MT01-MT10 composition in a city profile.", "", "Step09"),
        ("population_share_MTxx", "morphotype", "Population share of a local morphotype within a drive hex_1km city profile.", "Population share for one local morphotype in a drive hex_1km city profile.", "0-1", "Step13"),
        ("unit_share_MTxx", "morphotype", "Spatial-unit share of a local morphotype within a drive hex_1km city profile.", "Spatial-unit share for one local morphotype in a drive hex_1km city profile.", "0-1", "Step13"),
        ("population_share_total", "morphotype", "Population-weighted total share of a local morphotype across available drive hex_1km local units.", "Population-weighted total share across available drive hex_1km local units.", "0-1", "Step13"),
        ("unit_share_total", "morphotype", "Total spatial-unit share of a local morphotype across available drive hex_1km local units.", "Total spatial-unit share across available drive hex_1km local units.", "0-1", "Step13"),
        ("edge_density_km_per_km2", "network_metric", "Road edge length density.", "Road edge length density.", "km/km^2", "Step06"),
        ("intersection_density_per_km2", "network_metric", "Intersection density.", "Intersection density.", "count/km^2", "Step06"),
        ("four_way_share", "network_metric", "Share of four-way intersections.", "Share of four-way intersections.", "0-1", "Step06"),
        ("dead_end_share", "network_metric", "Share of dead-end nodes.", "Share of dead-end nodes.", "0-1", "Step06"),
        ("segment_length_median", "network_metric", "Median road segment length.", "Median road segment length.", "m", "Step06"),
        ("edge_circuity_mean", "network_metric", "Mean edge circuity.", "Mean road-edge circuity.", "ratio", "Step06"),
        ("orientation_order", "network_metric", "Street orientation order metric.", "Street orientation order metric.", "0-1", "Step06"),
        ("road_hierarchy_entropy", "network_metric", "Road hierarchy entropy.", "Road hierarchy entropy.", "", "Step06"),
        ("population_weighted_access_share", "accessibility", "Population-weighted share with access within threshold.", "Population-weighted reachable share.", "0-1", "Step10/11"),
        ("accessibility_gini", "accessibility", "Gini of accessibility within city/type/category.", "Accessibility Gini within city, type, or category.", "0-1", "Step10/11"),
        ("no_access_population_share", "accessibility", "Population-weighted no-access share.", "Population-weighted no-access share.", "0-1", "Step10/11"),
        ("od_circuity_weighted_mean", "detour", "Weighted mean OD route circuity.", "Weighted mean OD route circuity.", "ratio", "Step10/11"),
        ("detour_ratio_median", "detour", "Median route detour ratio.", "Median route detour ratio.", "ratio", "Step10/11"),
        ("route_found_weighted_share", "detour", "Population/OD-weighted route-found share.", "Population- or OD-weighted route-found share.", "0-1", "Step10/11"),
        ("coef", "model", "Regression coefficient.", "Regression coefficient.", "", "Step11"),
        ("std_err", "model", "Standard error.", "Standard error.", "", "Step11"),
        ("p_value", "model", "Model p-value.", "Model p-value.", "", "Step11"),
        ("ci_low", "model", "Lower confidence interval bound.", "Lower confidence interval bound.", "", "Step11"),
        ("ci_high", "model", "Upper confidence interval bound.", "Upper confidence interval bound.", "", "Step11"),
        ("nobs", "model", "Number of observations used by fitted model.", "Number of observations used by the fitted model.", "count", "Step11"),
        ("scenario", "robustness", "Main or robustness scenario.", "Main or robustness scenario label.", "", "Step11"),
        ("hash_sha256", "provenance", "SHA-256 file hash recorded by Step13.", "SHA-256 file hash recorded by Step13.", "", "Step13"),
        ("mtime", "provenance", "File modification time recorded by Step13.", "File modification time recorded by Step13.", "", "Step13"),
    ]
    return [
        {
            "variable_name": name,
            "variable_group": group,
            "definition_en": en,
            "definition_detail": detail,
            "unit": unit,
            "source_step": step,
        }
        for name, group, en, detail, unit, step in specs
    ]


def source_table_columns(paths: StepPaths) -> pd.DataFrame:
    specs = [
        ("city_master", paths.city_master),
        ("city_quality", paths.city_quality),
        ("type_centers", paths.type_centers),
        ("city_type_profiles", paths.city_type_profiles),
        ("od_metrics", paths.od_metrics),
        ("fig1_global_sample_cities_quality_tiers", paths.step12_fig1_data),
        ("fig2_osm_quality_confidence", paths.step12_fig2_data),
        ("fig3_multiscale_network_signatures", paths.step12_fig3_data),
        ("fig4_morphotype_atlas_signatures", paths.step12_fig4_data),
        ("fig5_morphotype_radar", paths.step12_fig5_data),
        ("fig6_scale_network_transition_matrix", paths.step12_fig6_data),
        ("fig7_core_city_overview", paths.step12_fig7_data),
        ("fig8_access_detour_performance", paths.step12_fig8_data),
        ("fig9_observed_group_means", paths.step12_fig9_data),
        ("fig10_city_inequality", paths.step12_fig10_data),
        ("fig11_model_coefficient_forest", paths.step12_fig11_data),
        ("city_inequality_source", paths.city_inequality_source),
        ("model_coefficients", paths.model_coefficients),
        ("robustness_results", paths.robustness_results),
    ]
    rows: list[dict[str, Any]] = []
    for table_id, path in specs:
        if not path.exists():
            continue
        if path.suffix.lower() == ".parquet":
            schema = pq.ParquetFile(path).schema_arrow
            names = schema.names
            dtypes = [str(schema.field(i).type) for i in range(len(names))]
        else:
            sample = pd.read_csv(path, nrows=20)
            names = list(sample.columns)
            dtypes = [str(sample[c].dtype) for c in names]
        for idx, (name, dtype) in enumerate(zip(names, dtypes), start=1):
            rows.append(
                {
                    "table_id": table_id,
                    "column_order": idx,
                    "variable_name": name,
                    "dtype": dtype,
                    "source_file": safe_rel(path, paths.root),
                }
            )
    return pd.DataFrame(rows)


def excel_safe_df(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out = out.replace([np.inf, -np.inf], np.nan)
    for col in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[col]):
            out[col] = out[col].astype(str)
        elif out[col].dtype == object:
            out[col] = out[col].map(
                lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list, tuple, set)) else v
            )
    return out


def sanitize_sheet_name(name: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]\:\*\?\/\\]", "_", name).strip() or "Sheet"
    clean = clean[:31]
    base = clean
    i = 1
    while clean in used:
        suffix = f"_{i}"
        clean = f"{base[:31 - len(suffix)]}{suffix}"
        i += 1
    used.add(clean)
    return clean


def write_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if len(sheets) < 2:
        raise ValueError(f"{path.name} must contain at least two sheets.")
    used: set[str] = set()
    sheet_map = {sanitize_sheet_name(name, used): excel_safe_df(df) for name, df in sheets.items()}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in sheet_map.items():
            if df.empty:
                raise ValueError(f"{path.name} sheet {sheet_name} is empty.")
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.book[sheet_name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            header_fill = PatternFill("solid", fgColor="D9EAF7")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col_idx, column_cells in enumerate(ws.columns, start=1):
                values = [str(c.value) for c in list(column_cells)[:200] if c.value is not None]
                width = min(max([len(v) for v in values] + [10]) + 2, 48)
                ws.column_dimensions[get_column_letter(col_idx)].width = width
            for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 500)):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=False, vertical="top")


def build_appendix_notes(title: str, notes: list[str]) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "note_id": [f"N{i:02d}" for i in range(1, len(notes) + 1)],
            "appendix": title,
            "note": notes,
        }
    )


def collect_input_records(paths: StepPaths) -> list[dict[str, Any]]:
    specs = [
        ("Step01", "city sample master", paths.city_master, True),
        ("Step01", "city boundaries", paths.city_boundaries, True),
        ("Step06", "city morphology metrics", paths.city_morphology_metrics, True),
        ("Step06", "local morphology metrics", paths.local_morphology_metrics, True),
        ("Step07", "city quality scores", paths.city_quality, True),
        ("Step07", "quality review checklist", paths.quality_review, False),
        ("Step08", "city profile base", paths.city_profile_base, True),
        ("Step08", "scale signature long table", paths.scale_signature_long, True),
        ("Step09", "type centers", paths.type_centers, True),
        ("Step09", "feature contributions", paths.type_contrib, False),
        ("Step09", "city type profiles", paths.city_type_profiles, True),
        ("Step09", "local morphotypes", paths.local_morphotypes, True),
        ("Step10", "OD detour metrics", paths.od_metrics, True),
        ("Step10", "OD detour samples", paths.od_samples, True),
        ("Step10", "facility accessibility metrics", paths.accessibility_metrics, True),
        ("Step10", "accessibility inequality metrics", paths.inequality_metrics, True),
        ("Step10", "facility classification table", paths.facility_classification, True),
        ("Step10", "low confidence sensitivity", paths.low_confidence_sensitivity, False),
        ("Step10", "Step10 QC", paths.step10_qc, True),
        ("Step11", "model A table", paths.model_a_table, True),
        ("Step11", "model B table", paths.model_b_table, True),
        ("Step11", "model C table", paths.model_c_table, True),
        ("Step11", "model coefficients", paths.model_coefficients, True),
        ("Step11", "model results", paths.model_results, True),
        ("Step11", "robustness results", paths.robustness_results, True),
        ("Step11", "robustness summary", paths.robustness_summary, True),
        ("Step11", "model marginal summary", paths.marginal_effects, True),
        ("Step11", "accessibility source data feeding Fig6C/Fig8", paths.fig7_source, True),
        ("Step11", "detour/circuity source data feeding Fig8", paths.fig8_source, True),
        ("Step11", "city inequality source data", paths.city_inequality_source, True),
        ("Step11", "Step11 QC", paths.step11_qc, True),
        ("Step12", "Step12 README metadata", paths.step12_readme, True),
        ("Step12", "Step12 execution log metadata", paths.step12_run_log, True),
        ("Step12", "Step12 reproducibility metadata", paths.step12_repro, True),
        ("Step12", "Step12 figure manifest", paths.step12_figure_manifest, True),
        ("Step12", "Step12 figure lineage", paths.step12_lineage, True),
        ("Step12", "Step12 Fig1 global sample cities and quality tiers data", paths.step12_fig1_data, True),
        ("Step12", "Step12 Fig2 OSM quality confidence data", paths.step12_fig2_data, True),
        ("Step12", "Step12 Fig3 multiscale network signatures data", paths.step12_fig3_data, True),
        ("Step12", "Step12 Fig4 morphotype atlas signatures data", paths.step12_fig4_data, True),
        ("Step12", "Step12 Fig5 morphotype radar signatures data", paths.step12_fig5_data, True),
        ("Step12", "Step12 Fig6 scale/network/performance heatmap data", paths.step12_fig6_data, True),
        ("Step12", "Step12 Fig7 core-city overview data", paths.step12_fig7_data, True),
        ("Step12", "Step12 Fig8 current A/B accessibility/detour performance data", paths.step12_fig8_data, True),
        ("Step12", "Step12 Fig9 observed weighted group means data", paths.step12_fig9_data, True),
        ("Step12", "Step12 Fig10 city-level accessibility inequality data", paths.step12_fig10_data, True),
        ("Step12", "Step12 Fig11 model coefficient forest data", paths.step12_fig11_data, True),
        ("Step12", "Step12 QC", paths.step12_qc, True),
        ("manuscript", "Table 1 maintained in manuscript", paths.manuscript, False),
    ]
    records: list[dict[str, Any]] = []
    for source_step, role, path, required in specs:
        rec = file_metadata(path, paths.root, artifact_role=role, source_step=source_step)
        rec["required_for_step13"] = required
        rec["input_role"] = role
        records.append(rec)
    return records


def input_meta_map(records: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {str(r["relative_path"]): r for r in records}


def write_input_report(paths: StepPaths, input_records: list[dict[str, Any]], protected_before_digest: str) -> dict[str, Any]:
    required_missing = [r for r in input_records if r.get("required_for_step13") and not r.get("exists")]
    report = {
        "step": STEP_NAME,
        "created_at": now_iso(),
        "root": str(paths.root),
        "output_dir": str(paths.output_dir),
        "required_input_count": sum(1 for r in input_records if r.get("required_for_step13")),
        "required_missing_count": len(required_missing),
        "required_missing": [r["relative_path"] for r in required_missing],
        "protected_upstream_manifest_before_sha256": protected_before_digest,
        "risk_notes": RISK_NOTES,
        "input_records": input_records,
    }
    write_json(paths.output_dir / "step13_input_acceptance_report.json", report)

    lines = [
        "# Step13 Input Acceptance Report",
        "",
        f"- Generated at: {report['created_at']}",
        f"- Project root: `{paths.root}`",
        f"- Output directory: `{paths.output_dir}`",
        f"- Required input count: {report['required_input_count']}",
        f"- Missing required input count: {report['required_missing_count']}",
        "",
        "## Non-blocking Interpretation Notes",
        "",
        *[f"- {note}" for note in RISK_NOTES],
        "",
        "## Input Files",
        "",
        "| source_step | input_role | relative_path | rows | columns | required | exists |",
        "|---|---|---:|---:|---:|---:|---:|",
    ]
    for r in input_records:
        lines.append(
            f"| {r.get('source_step','')} | {r.get('input_role','')} | `{r.get('relative_path','')}` | "
            f"{r.get('rows','')} | {r.get('columns','')} | {r.get('required_for_step13','')} | {r.get('exists','')} |"
        )
    (paths.output_dir / "step13_input_acceptance_report.md").write_text("\n".join(lines), encoding="utf-8")
    return report


def build_quality_warning_sheet(paths: StepPaths) -> pd.DataFrame:
    rows = []
    for step, qc_path in [("Step10", paths.step10_qc), ("Step11", paths.step11_qc), ("Step12", paths.step12_qc)]:
        if not qc_path.exists():
            continue
        df = pd.read_csv(qc_path)
        status_col = "status" if "status" in df.columns else None
        check_col = "check" if "check" in df.columns else "check_name" if "check_name" in df.columns else None
        if not status_col or not check_col:
            continue
        warn_fail = df[df[status_col].astype(str).str.lower().isin(["warn", "warning", "fail"])].copy()
        for _, r in warn_fail.iterrows():
            status = str(r.get(status_col, "")).lower()
            rows.append(
                {
                    "source_step": step,
                    "check": r.get(check_col),
                    "status": "warn" if status in {"warn", "warning"} else status,
                    "value": r.get("value", r.get("actual", "")),
                    "expected": r.get("expected", ""),
                    "detail": r.get("detail", ""),
                    "step13_policy": "non_blocking_warn_recorded"
                    if status in {"warn", "warning"}
                    else "upstream_fail_would_block",
                }
            )
    if not rows:
        rows.append(
            {
                "source_step": "Step13",
                "check": "upstream_warnings",
                "status": "pass",
                "value": 0,
                "expected": 0,
                "detail": "No upstream warnings/failures found in checked QC files.",
                "step13_policy": "recorded",
            }
        )
    out = pd.DataFrame(rows)
    out.loc[
        out["check"].astype(str).str.contains("small_groups|small-group|small_group", case=False, regex=True, na=False),
        "step13_policy",
    ] = "Step11 small-group warning is non-blocking and retained as a caveat."
    return out


def build_material_rows(
    paths: StepPaths,
    output_paths: list[Path],
    input_records: list[dict[str, Any]],
    include_self_placeholders: bool = True,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for p in output_paths:
        role = (
            "main_text_table"
            if p.name in MAIN_TABLE_FILES
            else "appendix_table"
            if p.name in APPENDIX_TABLE_FILES
            else "appendix_workbook"
            if p.name in APPENDIX_FILES
            else "step13_record"
        )
        if include_self_placeholders and p.name in {"paper_materials_manifest.csv", "paper_materials_manifest.xlsx"}:
            rows.append(
                {
                    "artifact_name": p.name,
                    "artifact_role": role,
                    "relative_path": safe_rel(p, paths.root),
                    "absolute_path": str(p.resolve()),
                    "exists": p.exists(),
                    "rows": None,
                    "columns": None,
                    "sheet_count": None,
                    "sheet_summary": "",
                    "size_bytes": p.stat().st_size if p.exists() else None,
                    "mtime": datetime.fromtimestamp(p.stat().st_mtime).astimezone().isoformat(timespec="seconds")
                    if p.exists()
                    else None,
                    "hash_sha256": "",
                    "source_step": "Step13",
                    "source_file": "self-referential material list; hash omitted to keep the inventory stable",
                    "material_note": "Step13 generated artifact.",
                }
            )
            continue
        rec = file_metadata(p, paths.root, artifact_role=role, source_step="Step13")
        rec["material_note"] = "Step13 generated artifact."
        rows.append(rec)

    figure_manifest = pd.read_csv(paths.step12_figure_manifest) if paths.step12_figure_manifest.exists() else pd.DataFrame()
    if not figure_manifest.empty:
        for _, row in figure_manifest.iterrows():
            for col, role in [
                ("png_path", "step12_figure"),
                ("svg_path", "step12_figure"),
                ("pdf_path", "step12_figure"),
                ("data_path", "step12_figure_data"),
            ]:
                value = row.get(col)
                if not isinstance(value, str) or not value:
                    continue
                p = Path(value)
                if not p.is_absolute():
                    p = paths.root / value
                rec = file_metadata(p, paths.root, artifact_role=role, source_step="Step12")
                rec["source_file"] = str(row.get("source_inputs", ""))
                figure_title = row.get("figure_title", row.get("title", ""))
                rec["material_note"] = f"{row.get('figure_id','')} {figure_title}"
                rows.append(rec)

    for r in input_records:
        rr = dict(r)
        rr["artifact_role"] = "key_upstream_input"
        rr["material_note"] = rr.get("input_role", "")
        rows.append(rr)

    rows.append(
        {
            "artifact_name": "table1_maintained_by_manuscript",
            "artifact_role": "manuscript_maintained_table_note",
            "relative_path": safe_rel(paths.manuscript, paths.root),
            "absolute_path": str(paths.manuscript),
            "exists": paths.manuscript.exists(),
            "rows": None,
            "columns": None,
            "sheet_count": None,
            "sheet_summary": "",
            "size_bytes": paths.manuscript.stat().st_size if paths.manuscript.exists() else None,
            "mtime": datetime.fromtimestamp(paths.manuscript.stat().st_mtime).astimezone().isoformat(timespec="seconds")
            if paths.manuscript.exists()
            else None,
            "hash_sha256": sha256_file(paths.manuscript) if paths.manuscript.exists() else None,
            "source_step": "manuscript",
            "source_file": "Table 1 is maintained in docs/manuscript.docx; Step13 does not modify manuscript files.",
            "material_note": "Do not auto-update Table 1 from Step13.",
        }
    )
    out = pd.DataFrame(rows)
    dedupe_cols = ["artifact_role", "relative_path", "material_note"]
    out = out.drop_duplicates(subset=[c for c in dedupe_cols if c in out.columns]).reset_index(drop=True)
    return out


def write_material_list(paths: StepPaths, material: pd.DataFrame) -> None:
    write_csv(material, paths.output_dir / "paper_materials_manifest.csv")
    dictionary = pd.DataFrame(
        [
            {"field": c, "definition": "Step13 material checklist field."}
            for c in material.columns
        ]
    )
    write_excel(
        paths.output_dir / "paper_materials_manifest.xlsx",
        {
            "material_checklist": material,
            "field_dictionary": dictionary,
        },
    )


def share_consistency_record(
    df: pd.DataFrame,
    morphotype_col: str,
    share_col: str,
    metric: str,
    *,
    max_examples: int = 5,
) -> dict[str, Any]:
    missing = [c for c in [morphotype_col, share_col] if c not in df.columns]
    matrix = morphotype_share_matrix(df, metric)
    if matrix.empty:
        missing.append(f"{metric}_share_MT01-MT10")
    if missing:
        return {
            "morphotype_col": morphotype_col,
            "share_col": share_col,
            "metric": metric,
            "missing": missing,
            "mismatch_count": None,
            "examples": [],
        }

    morphotypes = df[morphotype_col].astype("object")
    expected = lookup_morphotype_share(matrix, morphotypes)
    actual = pd.to_numeric(df[share_col], errors="coerce")
    ok = pd.Series(
        np.isclose(actual.to_numpy(dtype=float), expected.to_numpy(dtype=float), rtol=1e-9, atol=1e-12, equal_nan=True),
        index=df.index,
    )
    mismatches = df.loc[~ok]
    share_cols = morphotype_share_columns(df, metric)
    examples: list[dict[str, Any]] = []
    for idx, row in mismatches.head(max_examples).iterrows():
        mt = str(row.get(morphotype_col, ""))
        examples.append(
            {
                "row_index": int(idx) if isinstance(idx, (int, np.integer)) else str(idx),
                "city_id": row.get("city_id", ""),
                "morphotype": mt,
                "actual_share": actual.loc[idx],
                "expected_share": expected.loc[idx],
                "expected_column": share_cols.get(mt, ""),
            }
        )
    return {
        "morphotype_col": morphotype_col,
        "share_col": share_col,
        "metric": metric,
        "missing": [],
        "mismatch_count": int((~ok).sum()),
        "examples": examples,
    }


def column_equality_record(df: pd.DataFrame, left: str, right: str, *, numeric: bool = False, max_examples: int = 5) -> dict[str, Any]:
    missing = [c for c in [left, right] if c not in df.columns]
    if missing:
        return {"left": left, "right": right, "missing": missing, "mismatch_count": None, "examples": []}
    if numeric:
        left_values = pd.to_numeric(df[left], errors="coerce")
        right_values = pd.to_numeric(df[right], errors="coerce")
        ok = pd.Series(
            np.isclose(left_values.to_numpy(dtype=float), right_values.to_numpy(dtype=float), rtol=1e-9, atol=1e-12, equal_nan=True),
            index=df.index,
        )
    else:
        left_values = df[left].astype("string").fillna("")
        right_values = df[right].astype("string").fillna("")
        ok = left_values.eq(right_values)
    mismatches = df.loc[~ok]
    examples = [
        {
            "row_index": int(idx) if isinstance(idx, (int, np.integer)) else str(idx),
            "city_id": row.get("city_id", ""),
            "left_value": row.get(left, ""),
            "right_value": row.get(right, ""),
        }
        for idx, row in mismatches.head(max_examples).iterrows()
    ]
    return {"left": left, "right": right, "missing": [], "mismatch_count": int((~ok).sum()), "examples": examples}


def dominant_consistency_summary(df: pd.DataFrame) -> dict[str, Any]:
    share_checks = [
        share_consistency_record(df, "dominant_morphotype", "dominant_population_share", "population"),
        share_consistency_record(df, "dominant_morphotype", "dominant_unit_share", "unit"),
        share_consistency_record(df, "population_dominant_morphotype", "population_dominant_share", "population"),
        share_consistency_record(df, "unit_dominant_morphotype", "unit_dominant_share", "unit"),
    ]
    equality_checks = [
        column_equality_record(df, "dominant_morphotype", "population_dominant_morphotype"),
        column_equality_record(df, "dominant_morphotype_name", "population_dominant_morphotype_name"),
        column_equality_record(df, "dominant_population_share", "population_dominant_share", numeric=True),
    ]
    missing = [item for check in share_checks + equality_checks for item in check.get("missing", [])]
    mismatch_total = sum(int(check.get("mismatch_count") or 0) for check in share_checks + equality_checks)
    return {
        "rows": len(df),
        "missing": missing,
        "mismatch_total": mismatch_total,
        "share_checks": share_checks,
        "equality_checks": equality_checks,
    }


def build_qc(
    paths: StepPaths,
    tables: dict[str, pd.DataFrame],
    protected_changed: list[str],
    material: pd.DataFrame,
) -> pd.DataFrame:
    checks: list[dict[str, Any]] = []

    def add(check: str, status: str, value: Any, expected: Any, detail: str = "") -> None:
        checks.append({"check": check, "status": status, "value": value, "expected": expected, "detail": detail})

    required_paths = [paths.output_dir / name for name in OUTPUT_FILES]
    missing = [p.name for p in required_paths if not p.exists()]
    add("required_outputs_exist", "pass" if not missing else "fail", len(missing), 0, "; ".join(missing[:10]))

    main_nonempty = []
    for name in MAIN_TABLE_FILES:
        p = paths.output_dir / name
        rows, _ = csv_shape(p) if p.exists() else (0, 0)
        main_nonempty.append(rows is not None and rows > 0)
    add(
        f"main_text_csv_{len(MAIN_TABLE_FILES)}_exist_nonempty",
        "pass" if all(main_nonempty) else "fail",
        sum(main_nonempty),
        len(MAIN_TABLE_FILES),
    )

    appendix_ok = []
    appendix_details = []
    for name in APPENDIX_FILES:
        p = paths.output_dir / name
        rows, cols, sheets, summary = xlsx_shape(p) if p.exists() else (0, 0, 0, "")
        ok = bool(p.exists() and sheets and sheets >= 2 and rows and rows > sheets and cols and cols > 0)
        appendix_ok.append(ok)
        appendix_details.append(f"{name}:{summary}")
    add("appendix_A_to_F_xlsx_nonempty_min_2_sheets", "pass" if all(appendix_ok) else "fail", sum(appendix_ok), 6, " | ".join(appendix_details))

    role_col = material.get("artifact_role", pd.Series(dtype=str)).astype(str)
    roles = set(role_col)
    for role in [
        "main_text_table",
        "appendix_table",
        "appendix_workbook",
        "step12_figure",
        "step12_figure_data",
        "key_upstream_input",
    ]:
        add(f"material_checklist_contains_{role}", "pass" if role in roles else "fail", role in roles, True)

    step12_metadata_files = [
        paths.step12_readme,
        paths.step12_run_log,
        paths.step12_repro,
        paths.step12_qc,
        paths.step12_figure_manifest,
        paths.step12_lineage,
    ]
    metadata_expected = [safe_rel(p, paths.root) for p in step12_metadata_files]
    relative_col = material.get("relative_path", pd.Series(dtype=str)).astype(str)
    exists_col = material.get("exists", pd.Series(False, index=material.index)).fillna(False).astype(bool)
    metadata_present = set(relative_col[relative_col.isin(metadata_expected) & exists_col])
    metadata_missing = [p for p in metadata_expected if p not in metadata_present]
    add(
        "step12_metadata_files_in_material_manifest",
        "pass" if not metadata_missing else "fail",
        len(metadata_expected) - len(metadata_missing),
        len(metadata_expected),
        "; ".join(metadata_missing),
    )

    note_col = material.get("material_note", pd.Series(dtype=str)).astype(str).str.strip()
    hash_col = material.get("hash_sha256", pd.Series(dtype=str)).astype(str)
    name_col = material.get("artifact_name", pd.Series(dtype=str)).astype(str)

    new_table_names = {"table9_morphotype_definitions_and_representative_cities.csv", "city_morphotype_composition_full.csv"}
    new_table_material = material[name_col.isin(new_table_names)].copy()
    new_table_rows = pd.to_numeric(new_table_material.get("rows", pd.Series(dtype=float)), errors="coerce")
    new_table_hashes = new_table_material.get("hash_sha256", pd.Series(dtype=str)).astype(str)
    new_table_material_ok = (
        set(new_table_material.get("artifact_name", pd.Series(dtype=str)).astype(str)) == new_table_names
        and new_table_material.get("exists", pd.Series(False, index=new_table_material.index)).fillna(False).astype(bool).all()
        and new_table_rows.fillna(0).gt(0).all()
        and new_table_hashes.str.len().gt(0).all()
    )
    add(
        "material_checklist_contains_new_morphotype_tables",
        "pass" if new_table_material_ok else "fail",
        json.dumps(
            {
                "artifact_names": sorted(new_table_material.get("artifact_name", pd.Series(dtype=str)).astype(str).tolist()),
                "rows": new_table_rows.astype("Int64").astype(str).tolist(),
                "hash_present": new_table_hashes.str.len().gt(0).tolist(),
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Table 9 and city morphotype composition CSV are present, non-empty, and hashed in the material checklist",
    )

    def step12_material_hash_check(
        figure_id: str,
        expected_names: set[str],
        expected_detail: str,
    ) -> None:
        figure_material = material[
            note_col.str.startswith(figure_id)
            & role_col.isin(["step12_figure", "step12_figure_data"])
            & name_col.isin(expected_names)
        ].copy()
        hash_records: list[dict[str, Any]] = []
        mismatches: list[str] = []
        for idx, row in figure_material.iterrows():
            rel_value = row.get("relative_path", "")
            path = Path(str(row.get("absolute_path", "")))
            if not path.is_absolute():
                path = paths.root / str(rel_value)
            current_hash = sha256_file(path)
            manifest_hash = str(hash_col.loc[idx])
            matches = bool(current_hash and manifest_hash and current_hash == manifest_hash)
            if not matches:
                mismatches.append(str(row.get("artifact_name", "")))
            hash_records.append(
                {
                    "artifact_name": row.get("artifact_name"),
                    "relative_path": rel_value,
                    "manifest_hash": manifest_hash,
                    "current_hash": current_hash,
                    "matches": matches,
                }
            )
        missing_material = sorted(expected_names - set(figure_material.get("artifact_name", pd.Series(dtype=str)).astype(str)))
        hashes_ok = bool(not missing_material and not mismatches and len(figure_material) == len(expected_names))
        add(
            f"{figure_id}_step13_material_hashes_match_current_files",
            "pass" if hashes_ok else "fail",
            json.dumps(
                {
                    "records": hash_records,
                    "missing_material_rows": missing_material,
                    "mismatches": mismatches,
                },
                ensure_ascii=False,
                sort_keys=True,
            ),
            expected_detail,
        )

    canonical_figure_stems = {
        "Fig1": "Fig1_global_sample_cities_quality_tiers",
        "Fig2": "Fig2_osm_quality_confidence",
        "Fig3": "Fig3_multiscale_network_signatures",
        "Fig4": "Fig4_morphotype_atlas_signatures",
        "Fig5": "Fig5_morphotype_radar_signatures",
        "Fig6": "Fig6_scale_network_transition_matrix",
        "Fig7": "Fig7_core_city_morphotype_stability_and_performance",
        "Fig8": "Fig8_accessibility_detour_performance_by_morphotype",
        "Fig9": "Fig9_observed_weighted_group_means",
        "Fig10": "Fig10_city_level_accessibility_inequality",
        "Fig11": "Fig11_model_coefficient_forest",
    }
    for fig_id, stem in canonical_figure_stems.items():
        step12_material_hash_check(
            fig_id,
            {
                f"{stem}.png",
                f"{stem}.svg",
                f"{stem}.pdf",
                f"figure_data_{stem}.csv",
            },
            f"Step13 material manifest {fig_id} PNG/SVG/PDF/source CSV hashes equal current Step12 files; TIFF exports are outside the Step13 material-list scope",
        )
    if paths.step12_figure_manifest.exists():
        step12_manifest = pd.read_csv(paths.step12_figure_manifest)
    else:
        step12_manifest = pd.DataFrame()
    canonical_figures = step12_manifest[
        step12_manifest.get("record_role", pd.Series(dtype=str)).astype(str).eq("canonical")
    ].copy()
    expected_all_fig_ids = {f"Fig{i}" for i in range(1, 12)}
    all_fig_expected: list[dict[str, str]] = []
    for _, row in canonical_figures.iterrows():
        fig_id = str(row.get("figure_id", ""))
        for col in ["png_path", "svg_path", "pdf_path", "data_path"]:
            value = row.get(col)
            if not isinstance(value, str) or not value:
                continue
            p = Path(value)
            if not p.is_absolute():
                p = paths.root / value
            all_fig_expected.append({"figure_id": fig_id, "artifact_name": p.name, "path": str(p)})
    all_missing_material: list[str] = []
    all_mismatches: list[str] = []
    all_hash_records: list[dict[str, Any]] = []
    for rec in all_fig_expected:
        matches = material[
            role_col.isin(["step12_figure", "step12_figure_data"])
            & name_col.eq(rec["artifact_name"])
            & note_col.str.startswith(rec["figure_id"])
        ].copy()
        if matches.empty:
            all_missing_material.append(f"{rec['figure_id']}:{rec['artifact_name']}")
            continue
        row = matches.iloc[0]
        current_hash = sha256_file(Path(rec["path"]))
        manifest_hash = str(row.get("hash_sha256", ""))
        ok = bool(current_hash and manifest_hash and current_hash == manifest_hash)
        if not ok:
            all_mismatches.append(f"{rec['figure_id']}:{rec['artifact_name']}")
        all_hash_records.append(
            {
                "figure_id": rec["figure_id"],
                "artifact_name": rec["artifact_name"],
                "current_hash": current_hash,
                "manifest_hash": manifest_hash,
                "matches": ok,
            }
        )
    canonical_ids = set(canonical_figures.get("figure_id", pd.Series(dtype=str)).astype(str))
    all_fig_sync_ok = (
        canonical_ids == expected_all_fig_ids
        and len(all_fig_expected) == 44
        and not all_missing_material
        and not all_mismatches
        and len(all_hash_records) == 44
    )
    add(
        "Fig1_Fig11_step13_material_hashes_match_current_files",
        "pass" if all_fig_sync_ok else "fail",
        json.dumps(
            {
                "canonical_ids": sorted(canonical_ids),
                "expected_artifact_count": 44,
                "observed_artifact_count_from_step12_manifest": len(all_fig_expected),
                "checked_artifact_count": len(all_hash_records),
                "missing_material_rows": all_missing_material,
                "mismatches": all_mismatches,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "All 11 Step12 canonical main figures have PNG/SVG/PDF/source CSV rows in Step13 material manifest with hashes equal to current files; TIFF exports are intentionally not listed in Step13 materials",
    )

    table2_total = int(tables["table2"].loc[tables["table2"]["region"].eq("ALL"), "total_city_count"].iloc[0])
    table2_core = int(tables["table2"].loc[tables["table2"]["region"].eq("ALL"), "core_analysis_city_count"].iloc[0])
    add("city_count_86", "pass" if table2_total == 86 else "fail", table2_total, 86)
    add("main_analysis_core_city_count_48", "pass" if table2_core == 48 else "fail", table2_core, 48)

    type_count = int(tables["type_centers"]["morphotype"].nunique()) if "morphotype" in tables["type_centers"].columns else 0
    add("morphotype_count_10", "pass" if type_count == 10 else "fail", type_count, 10)

    table9 = tables["table9"]
    table9_required = {
        "morphotype",
        "morphotype_name_existing",
        "english_label",
        "structural_signature",
        "dominant_axes",
        "representative_core_cities",
        "n_dominant_core_cities",
        "population_share_total",
        "unit_share_total",
        "mean_type_confidence",
        "selection_rule",
    }
    table9_missing = sorted(table9_required - set(table9.columns))
    table9_mts = sorted(table9.get("morphotype", pd.Series(dtype=str)).astype(str).tolist())
    table9_ok = len(table9) == 10 and table9_mts == MORPHOTYPES and not table9_missing
    add(
        "table9_morphotype_definitions_10_rows_MT01_MT10",
        "pass" if table9_ok else "fail",
        json.dumps(
            {"rows": len(table9), "morphotypes": table9_mts, "missing_fields": table9_missing},
            ensure_ascii=False,
            sort_keys=True,
        ),
        "10 rows, MT01-MT10 complete, required fields present",
    )
    table9_rep = table9.get("representative_core_cities", pd.Series(dtype=str)).astype(str).str.strip()
    table9_note = table9.get("note", pd.Series(dtype=str)).astype(str).str.strip()
    n_dom = pd.to_numeric(table9.get("n_dominant_core_cities", pd.Series(dtype=float)), errors="coerce").fillna(0)
    representative_ok = table9_rep.ne("").all() and table9_note[n_dom.eq(0)].ne("").all()
    add(
        "table9_representative_core_cities_nonempty_or_noted",
        "pass" if representative_ok else "fail",
        json.dumps(
            {
                "empty_representative_morphotypes": table9.loc[table9_rep.eq(""), "morphotype"].astype(str).tolist(),
                "zero_dominant_without_note": table9.loc[n_dom.eq(0) & table9_note.eq(""), "morphotype"].astype(str).tolist(),
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Representative city list is non-empty; zero-dominant core types carry a note",
    )

    composition = tables["city_morphotype_composition"]
    reps = tables["representative_cities_by_type"]
    top5_mismatches: list[dict[str, Any]] = []
    core_for_reps = composition[
        composition.get("quality_tier", pd.Series(dtype=str)).astype(str).eq("core")
        & composition.get("network_type", pd.Series(dtype=str)).astype(str).eq("drive")
        & composition.get("scale", pd.Series(dtype=str)).astype(str).eq("hex_1km")
    ].copy()
    for mt in MORPHOTYPES:
        pop_col = f"population_share_{mt}"
        ranked = core_for_reps.copy()
        ranked["_mt_population_share"] = pd.to_numeric(ranked.get(pop_col), errors="coerce").fillna(0)
        expected_ids = (
            ranked.sort_values(["_mt_population_share", "city_name_en", "city_id"], ascending=[False, True, True])
            .head(5)["city_id"]
            .astype(str)
            .tolist()
        )
        actual_ids = (
            reps[reps.get("morphotype", pd.Series(dtype=str)).astype(str).eq(mt)]
            .sort_values("rank")["city_id"]
            .astype(str)
            .tolist()
        )
        if actual_ids != expected_ids:
            top5_mismatches.append({"morphotype": mt, "expected_city_ids": expected_ids, "actual_city_ids": actual_ids})
    add(
        "table9_representative_cities_population_share_top5_rule",
        "pass" if not top5_mismatches else "fail",
        json.dumps({"mismatches": top5_mismatches}, ensure_ascii=False, sort_keys=True),
        "Each MT representative list uses top five core drive hex_1km city profiles ranked by population_share_MTxx",
    )

    composition_required = {
        "city_id",
        "city_name_en",
        "country",
        "iso3",
        "region",
        "sample_group",
        "quality_tier",
        "network_type",
        "scale",
        "population_dominant_morphotype",
        "population_dominant_morphotype_name",
        "population_dominant_share",
        "unit_dominant_morphotype",
        "unit_dominant_morphotype_name",
        "unit_dominant_share",
        "dominant_morphotype",
        "dominant_morphotype_name",
        "dominant_population_share",
        "dominant_unit_share",
        "morphotype_entropy",
        "mean_type_confidence",
        "n_units",
    }
    composition_required.update({f"population_share_{mt}" for mt in MORPHOTYPES})
    composition_required.update({f"unit_share_{mt}" for mt in MORPHOTYPES})
    composition_missing = sorted(composition_required - set(composition.columns))
    source_profiles = tables["city_type_profiles"]
    expected_composition_rows = len(
        source_profiles[
            source_profiles["network_type"].astype(str).eq("drive")
            & source_profiles["scale"].astype(str).eq("hex_1km")
        ]
    )
    composition_drive_hex = (
        composition.get("network_type", pd.Series(dtype=str)).astype(str).eq("drive")
        & composition.get("scale", pd.Series(dtype=str)).astype(str).eq("hex_1km")
    )
    composition_ok = (
        len(composition) == expected_composition_rows
        and 80 <= len(composition) <= 90
        and composition_drive_hex.all()
        and not composition_missing
    )
    add(
        "city_morphotype_composition_drive_hex_1km_records_and_share_fields",
        "pass" if composition_ok else "fail",
        json.dumps(
            {
                "rows": len(composition),
                "expected_from_step09_city_profiles": expected_composition_rows,
                "all_drive_hex_1km": bool(composition_drive_hex.all()) if len(composition_drive_hex) else False,
                "missing_fields": composition_missing,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Drive hex_1km city profile rows match Step09 actual count (about 85/86) and MT share fields are present",
    )
    composition_dominant = dominant_consistency_summary(composition)
    add(
        "city_morphotype_composition_dominant_share_fields_consistent",
        "pass" if not composition_dominant["missing"] and composition_dominant["mismatch_total"] == 0 else "fail",
        json.dumps(json_ready(composition_dominant), ensure_ascii=False, sort_keys=True),
        "dominant_population_share equals population_share_{dominant_morphotype}; population_dominant_* and unit_dominant_* shares match their MT share columns",
    )

    appendix_d_sheets = set(xlsx_sheet_names(paths.output_dir / "appendix_d_clustering_and_type_naming.xlsx"))
    appendix_d_required = {"morphotype_definitions", "city_morphotype_composition", "representative_cities_by_type"}
    appendix_d_missing = sorted(appendix_d_required - appendix_d_sheets)
    add(
        "appendixD_contains_new_morphotype_sheets",
        "pass" if not appendix_d_missing else "fail",
        sorted(appendix_d_sheets),
        sorted(appendix_d_required),
        "; ".join(appendix_d_missing),
    )
    appendix_d_dominant_records: dict[str, Any] = {}
    appendix_d_dominant_ok = True
    for sheet in ["city_morphotype_composition", "city_type_profiles"]:
        if sheet not in appendix_d_sheets:
            appendix_d_dominant_records[sheet] = {"missing_sheet": True}
            appendix_d_dominant_ok = False
            continue
        try:
            sheet_df = pd.read_excel(paths.output_dir / "appendix_d_clustering_and_type_naming.xlsx", sheet_name=sheet)
            summary = dominant_consistency_summary(sheet_df)
            appendix_d_dominant_records[sheet] = summary
            appendix_d_dominant_ok = appendix_d_dominant_ok and not summary["missing"] and summary["mismatch_total"] == 0
        except Exception as exc:
            appendix_d_dominant_records[sheet] = {"read_error": str(exc)}
            appendix_d_dominant_ok = False
    add(
        "appendixD_city_profile_dominant_share_fields_consistent",
        "pass" if appendix_d_dominant_ok else "fail",
        json.dumps(json_ready(appendix_d_dominant_records), ensure_ascii=False, sort_keys=True),
        "Appendix D city_morphotype_composition and city_type_profiles sheets use the same population/unit dominant semantics",
    )

    def unique_nonempty(df: pd.DataFrame, col: str) -> list[str]:
        if col not in df.columns:
            return []
        return sorted(df[col].dropna().astype(str).str.strip().loc[lambda s: s.ne("")].unique().tolist())

    def unique_numbers(df: pd.DataFrame, col: str) -> list[float]:
        if col not in df.columns:
            return []
        vals = pd.to_numeric(df[col], errors="coerce").dropna().unique().tolist()
        return sorted(float(v) for v in vals)

    def bool_column_all(df: pd.DataFrame, col: str, expected: bool) -> bool:
        if col not in df.columns:
            return False
        vals = {str(v).strip().lower() for v in df[col].dropna().unique().tolist()}
        if expected:
            return vals == {"true"}
        return vals == {"false"}

    fig6 = tables["fig6_scale_step12"]
    fig6_required = {
        "panel",
        "morphotype",
        "fig7_panel_label",
        "mean_unit_share_percent",
        "transition_delta_pp",
        "access_share_diff_pp",
        "population_weighted_access_share",
        "category_mean_access_share",
        "fig7_panel_c_migration_note",
        "fig7_panel_c_norm_type",
        "fig7_panel_c_norm_vmin",
        "fig7_panel_c_norm_vcenter",
        "fig7_panel_c_norm_vmax",
        "fig7_panel_c_colorbar_label",
        "fig7_panel_c_color_variable",
    }
    fig6_missing = sorted(fig6_required - set(fig6.columns))
    fig6_panel_counts = fig6.get("panel", pd.Series(dtype=str)).astype(str).value_counts().to_dict()
    fig6_panel_labels = unique_nonempty(fig6, "fig7_panel_label")
    fig6_panel_c = fig6[fig6.get("panel", pd.Series(dtype=str)).astype(str).eq("performance_deviation")].copy()
    if {"population_weighted_access_share", "category_mean_access_share", "access_share_diff_pp"}.issubset(fig6_panel_c.columns):
        recomputed_diff = (
            pd.to_numeric(fig6_panel_c["population_weighted_access_share"], errors="coerce")
            - pd.to_numeric(fig6_panel_c["category_mean_access_share"], errors="coerce")
        ) * 100
        recorded_diff = pd.to_numeric(fig6_panel_c["access_share_diff_pp"], errors="coerce")
        fig6_recalc_error = float((recorded_diff - recomputed_diff).abs().max())
    else:
        fig6_recalc_error = None
    fig6_norm_checks = {
        "panel_c_norm_type_TwoSlopeNorm": unique_nonempty(fig6_panel_c, "fig7_panel_c_norm_type") == ["TwoSlopeNorm"],
        "panel_c_norm_vmin_-40": unique_numbers(fig6_panel_c, "fig7_panel_c_norm_vmin") == [-40.0],
        "panel_c_norm_vcenter_0": unique_numbers(fig6_panel_c, "fig7_panel_c_norm_vcenter") == [0.0],
        "panel_c_norm_vmax_40": unique_numbers(fig6_panel_c, "fig7_panel_c_norm_vmax") == [40.0],
        "panel_c_colorbar_label": unique_nonempty(fig6_panel_c, "fig7_panel_c_colorbar_label") == ["Difference from mean (pp)"],
        "panel_c_color_variable": unique_nonempty(fig6_panel_c, "fig7_panel_c_color_variable") == ["access_share_diff_pp"],
        "panel_c_migration_note": any("access-performance" in x and "Fig6C" in x for x in unique_nonempty(fig6_panel_c, "fig7_panel_c_migration_note")),
    }
    fig6_source_ok = (
        not fig6_missing
        and len(fig6) == 130
        and fig6_panel_counts.get("composition", 0) == 40
        and fig6_panel_counts.get("transition", 0) == 40
        and fig6_panel_counts.get("performance_deviation", 0) == 50
        and fig6_panel_labels == ["A", "B", "C"]
        and fig6_panel_c.get("morphotype", pd.Series(dtype=str)).nunique() == 10
        and fig6_panel_c.get("category", pd.Series(dtype=str)).nunique() == 5
        and fig6_recalc_error is not None
        and fig6_recalc_error <= 1e-9
        and all(fig6_norm_checks.values())
    )
    add(
        "Fig6_step12_A_B_C_migration_source_synced",
        "pass" if fig6_source_ok else "fail",
        json.dumps(
            {
                "rows": len(fig6),
                "panel_counts": fig6_panel_counts,
                "panel_labels": fig6_panel_labels,
                "panel_c_morphotype_count": int(fig6_panel_c.get("morphotype", pd.Series(dtype=str)).nunique()),
                "panel_c_category_count": int(fig6_panel_c.get("category", pd.Series(dtype=str)).nunique()),
                "access_share_diff_pp_recalculation_error": fig6_recalc_error,
                "missing_fields": fig6_missing,
                "norm_checks": fig6_norm_checks,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Fig6 source has A composition=40, B transition=40, C migrated access-performance deviation=50; Fig6C access_share_diff_pp and TwoSlopeNorm(-40,0,40) metadata are synced",
    )

    fig6_rel = safe_rel(paths.step12_fig6_data, paths.root)
    fig6_material = material[relative_col.eq(fig6_rel)]
    fig6_material_roles = set(role_col.loc[fig6_material.index])
    fig6_material_rows = pd.to_numeric(fig6_material.get("rows", pd.Series(dtype=float)), errors="coerce")
    fig6_material_ok = (
        {"step12_figure_data", "key_upstream_input"}.issubset(fig6_material_roles)
        and not fig6_material.empty
        and fig6_material_rows.eq(len(fig6)).all()
    )
    fig6_material_detail = "; ".join(
        f"{role}:rows={row}"
        for role, row in zip(
            role_col.loc[fig6_material.index].astype(str),
            fig6_material_rows.astype("Int64").astype(str),
        )
    )
    add(
        "Fig6_step12_csv_material_rows_match_source_csv",
        "pass" if fig6_material_ok else "fail",
        fig6_material_detail,
        f"step12_figure_data and key_upstream_input rows={len(fig6)}",
    )

    fig7 = tables["fig8_access_step12"]
    fig7_required = {
        "panel",
        "morphotype",
        "access_metric",
        "access_metric_name",
        "circuity_metric",
        "circuity_metric_name",
        "quadrant_label",
        "threshold_method",
        "p10_circuity",
        "p90_circuity",
        "fig9_panel_a_removed",
        "fig9_panel_a_heatmap_current",
        "fig9_panel_a_colorbar_current",
        "fig9_former_panel_a_migrated_to",
        "fig9_current_visual_panels",
        "fig9_layout_panel_b_y_axis_restored",
        "fig9_panel_c_marker_size_current",
        "fig9_panel_c_leader_line_count",
        "fig9_panel_c_label_offset_points",
    }
    fig7_missing = sorted(fig7_required - set(fig7.columns))
    fig7_panel_counts = fig7.get("panel", pd.Series(dtype=str)).astype(str).value_counts().to_dict()
    fig7_panel_labels = unique_nonempty(fig7, "fig9_panel_label")
    fig7_visual_panels = unique_nonempty(fig7, "fig9_current_visual_panels")
    fig7_marker_sizes = unique_numbers(fig7, "fig9_panel_c_marker_size_current")
    fig7_leader_counts = unique_numbers(fig7, "fig9_panel_c_leader_line_count")
    fig7_leader_label_values = unique_nonempty(fig7, "fig9_panel_c_leader_line_labels")
    fig7_leader_labels_from_source = (
        [x.strip() for x in fig7_leader_label_values[0].split(",") if x.strip()]
        if fig7_leader_label_values
        else []
    )

    def fig7_offset_for_morphotype(mt: str) -> Any:
        if "fig9_panel_c_label_offset_points" not in fig7.columns:
            return None
        vals = (
            fig7.loc[
                fig7.get("morphotype", pd.Series(dtype=str)).astype(str).eq(mt),
                "fig9_panel_c_label_offset_points",
            ]
            .dropna()
            .astype(str)
            .unique()
            .tolist()
        )
        parsed: list[Any] = []
        for value in vals:
            try:
                item = json.loads(value)
            except Exception:
                item = value
            if item not in parsed:
                parsed.append(item)
        return parsed[0] if len(parsed) == 1 else parsed

    fig7_source_offsets = {
        "MT03": fig7_offset_for_morphotype("MT03"),
        "MT04": fig7_offset_for_morphotype("MT04"),
    }
    fig7_ok = (
        not fig7_missing
        and len(fig7) == 20
        and fig7_panel_counts.get("access_panel", 0) == 0
        and fig7_panel_counts.get("detour_circuity_panel", 0) == 10
        and fig7_panel_counts.get("quadrant_panel", 0) == 10
        and fig7_panel_labels == ["A", "B"]
        and fig7_visual_panels == ["A, B"]
        and bool_column_all(fig7, "fig9_panel_a_removed", True)
        and bool_column_all(fig7, "fig9_panel_a_heatmap_current", False)
        and bool_column_all(fig7, "fig9_panel_a_colorbar_current", False)
        and unique_nonempty(fig7, "fig9_former_panel_a_migrated_to") == ["Fig6C"]
        and bool_column_all(fig7, "fig9_layout_panel_b_y_axis_restored", True)
        and fig7_marker_sizes == [30.0]
        and fig7_leader_counts == [8.0]
        and fig7_leader_labels_from_source == ["MT01", "MT03", "MT04", "MT05", "MT06", "MT07", "MT08", "MT09"]
        and fig7_source_offsets["MT03"] == [9, 5, "left"]
        and fig7_source_offsets["MT04"] == [-11, -3, "right"]
    )
    add(
        "Fig8_step12_A_B_only_required_fields_panel_counts",
        "pass" if fig7_ok else "fail",
        json.dumps(
            {
                "panel_counts": fig7_panel_counts,
                "panel_labels": fig7_panel_labels,
                "current_visual_panels": fig7_visual_panels,
                "panel_a_removed_all": bool_column_all(fig7, "fig9_panel_a_removed", True),
                "panel_a_heatmap_current_all_false": bool_column_all(fig7, "fig9_panel_a_heatmap_current", False),
                "panel_a_colorbar_current_all_false": bool_column_all(fig7, "fig9_panel_a_colorbar_current", False),
                "former_panel_a_migrated_to": unique_nonempty(fig7, "fig9_former_panel_a_migrated_to"),
                "panel_a_y_axis_restored_all": bool_column_all(fig7, "fig9_layout_panel_b_y_axis_restored", True),
                "panel_b_marker_sizes": fig7_marker_sizes,
                "panel_b_leader_counts": fig7_leader_counts,
                "panel_b_leader_labels": fig7_leader_labels_from_source,
                "panel_b_label_offsets": fig7_source_offsets,
                "missing_fields": fig7_missing,
                "rows": len(fig7),
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Current Fig8 Step12 CSV has no access_panel rows and exactly A/B visual panels: detour=10 and quadrant=10, with no current access heatmap/colorbar",
    )

    fig7_rel = safe_rel(paths.step12_fig8_data, paths.root)
    fig7_material = material[relative_col.eq(fig7_rel)]
    fig7_material_roles = set(role_col.loc[fig7_material.index])
    fig7_material_rows = pd.to_numeric(fig7_material.get("rows", pd.Series(dtype=float)), errors="coerce")
    fig7_material_ok = (
        {"step12_figure_data", "key_upstream_input"}.issubset(fig7_material_roles)
        and not fig7_material.empty
        and fig7_material_rows.eq(len(fig7)).all()
    )
    fig7_material_detail = "; ".join(
        f"{role}:rows={row}"
        for role, row in zip(
            role_col.loc[fig7_material.index].astype(str),
            fig7_material_rows.astype("Int64").astype(str),
        )
    )
    add(
        "Fig8_step12_csv_material_rows_match_combined_csv",
        "pass" if fig7_material_ok else "fail",
        fig7_material_detail,
        f"step12_figure_data and key_upstream_input rows={len(fig7)}",
    )

    step12_qc = tables["step12_qc"]
    step12_check_col = "check" if "check" in step12_qc.columns else "check_name" if "check_name" in step12_qc.columns else None
    step12_status_col = "status" if "status" in step12_qc.columns else None
    step12_value_col = "value" if "value" in step12_qc.columns else None

    def step12_qc_record(check_name: str) -> tuple[list[str], dict[str, Any], bool]:
        if not step12_check_col or not step12_status_col:
            return [], {}, False
        rows = step12_qc[step12_qc[step12_check_col].astype(str).eq(check_name)]
        statuses = rows[step12_status_col].astype(str).str.lower().tolist()
        record: dict[str, Any] = {}
        if not rows.empty and step12_value_col:
            raw_value = rows.iloc[0].get(step12_value_col, "")
            try:
                parsed = json.loads(str(raw_value))
                if isinstance(parsed, dict):
                    record = parsed
            except Exception:
                record = {"raw_value": raw_value}
        return statuses, record, bool(statuses and all(s == "pass" for s in statuses))

    def step12_qc_details(check_name: str) -> list[str]:
        if not step12_check_col or "detail" not in step12_qc.columns:
            return []
        rows = step12_qc[step12_qc[step12_check_col].astype(str).eq(check_name)]
        return rows["detail"].dropna().astype(str).tolist()

    expected_leader_labels = ["MT01", "MT03", "MT04", "MT05", "MT06", "MT07", "MT08", "MT09"]
    fig6c_statuses, fig6c_record, fig6c_status_ok = step12_qc_record(
        "Fig7_panel_C_migrated_access_performance_deviation_valid"
    )
    fig6c_record_counts = fig6c_record.get("panel_counts", {})
    fig6c_record_checks = {
        "step12_qc_status_pass": fig6c_status_ok,
        "panel_counts_match": fig6c_record_counts.get("composition") == 40
        and fig6c_record_counts.get("transition") == 40
        and fig6c_record_counts.get("performance_deviation") == 50,
        "panel_c_rows_50": fig6c_record.get("panel_c_rows") == 50,
        "panel_c_morphotype_count_10": fig6c_record.get("panel_c_morphotype_count") == 10,
        "panel_c_category_count_5": fig6c_record.get("panel_c_category_count") == 5,
        "recalculation_error_tiny": float(fig6c_record.get("max_diff_recalculation_error", 1.0)) <= 1e-9,
        "colorbar_label_synced": fig6c_record.get("panel_c_colorbar_labels") == ["Difference from mean (pp)"],
        "migration_note_synced": any("access-performance" in str(x) and "Fig6C" in str(x) for x in fig6c_record.get("panel_c_migration_notes", [])),
        "norm_record_synced": "TwoSlopeNorm(vmin=-40, vcenter=0, vmax=40)"
        in str(fig6c_record.get("panel_norms_by_panel", {}).get("performance_deviation", "")),
    }
    add(
        "Step12_QC_Fig6C_migrated_access_performance_synced",
        "pass" if all(fig6c_record_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig6c_statuses,
                "checks": fig6c_record_checks,
                "record": fig6c_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass for Fig6C migrated access-performance-deviation heatmap, including access_share_diff_pp and TwoSlopeNorm(-40,0,40)",
    )

    fig6_order_statuses, fig6_order_record, fig6_order_status_ok = step12_qc_record(
        "Fig7_panel_C_morphotype_order_aligned_with_A_B"
    )
    fig6_order_checks = {
        "step12_qc_status_pass": fig6_order_status_ok,
        "source_display_order_MT01_MT10": fig6_order_record.get("source_display_order") == MORPHOTYPES,
        "expected_order_MT01_MT10": fig6_order_record.get("expected_order") == MORPHOTYPES,
        "source_axis_order_MT01_to_MT10": fig6_order_record.get("source_axis_orders") == ["MT01->MT10"],
        "svg_tick_order_ok": fig6_order_record.get("svg_tick_order_ok") is True,
        "svg_old_axis_label_absent": fig6_order_record.get("svg_old_axis_label_absent") is True,
        "lineage_note_present": fig6_order_record.get("lineage_note_present") is True,
        "manifest_note_present": fig6_order_record.get("manifest_note_present") is True,
    }
    add(
        "Step12_QC_Fig6C_morphotype_order_MT01_MT10_synced",
        "pass" if all(fig6_order_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig6_order_statuses,
                "checks": fig6_order_checks,
                "record": fig6_order_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass for Fig6C morphotype order MT01->MT10 aligned with Fig6A/Fig6B; old ranked-by-access axis label absent",
    )

    fig6_text_statuses, fig6_text_record, fig6_text_status_ok = step12_qc_record(
        "Fig7_title_and_bottom_note_removed_from_render"
    )
    fig6_retained_terms = fig6_text_record.get("retained_terms_present", {})
    fig6_text_checks = {
        "step12_qc_status_pass": fig6_text_status_ok,
        "source_figure_level_title_current_false": fig6_text_record.get("source_figure_level_title_current_false") is True,
        "source_bottom_note_current_false": fig6_text_record.get("source_bottom_note_current_false") is True,
        "removed_title_absent_in_svg": fig6_text_record.get("removed_title_absent_in_svg") is True,
        "removed_bottom_note_absent_in_svg": fig6_text_record.get("removed_bottom_note_absent_in_svg") is True,
        "panel_c_old_axis_label_absent_in_svg": fig6_text_record.get("panel_c_old_axis_label_absent_in_svg") is True,
        "retained_terms_present": bool(fig6_retained_terms) and all(bool(v) for v in fig6_retained_terms.values()),
    }
    add(
        "Step12_QC_Fig6_title_bottom_note_removed_synced",
        "pass" if all(fig6_text_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig6_text_statuses,
                "checks": fig6_text_checks,
                "record": fig6_text_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass for Fig6 rendered canvas with no figure-level top title and no bottom footnote/note",
    )

    fig7_bc_statuses, fig7_bc_record, fig7_bc_status_ok = step12_qc_record(
        "Fig8_current_panels_A_B_only_no_access_heatmap_colorbar"
    )
    fig7_bc_counts = fig7_bc_record.get("panel_counts", {})
    fig7_bc_record_checks = {
        "step12_qc_status_pass": fig7_bc_status_ok,
        "access_panel_rows_0": fig7_bc_record.get("access_panel_rows") == 0,
        "detour_panel_rows_10": fig7_bc_record.get("detour_panel_rows") == 10,
        "quadrant_panel_rows_10": fig7_bc_record.get("quadrant_panel_rows") == 10,
        "panel_counts_match": fig7_bc_counts.get("detour_circuity_panel") == 10
        and fig7_bc_counts.get("quadrant_panel") == 10
        and fig7_bc_counts.get("access_panel", 0) == 0,
        "panel_labels_A_B": fig7_bc_record.get("panel_labels") == ["A", "B"],
        "panel_a_removed_all": fig7_bc_record.get("panel_a_removed_all") is True,
        "panel_a_heatmap_current_all_false": fig7_bc_record.get("panel_a_heatmap_current_all_false") is True,
        "panel_a_colorbar_current_all_false": fig7_bc_record.get("panel_a_colorbar_current_all_false") is True,
        "forbidden_a_current_fields_empty": fig7_bc_record.get("forbidden_a_current_fields_present", []) == [],
    }
    add(
        "Step12_QC_Fig8_A_B_only_no_access_heatmap_colorbar_synced",
        "pass" if all(fig7_bc_record_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig7_bc_statuses,
                "checks": fig7_bc_record_checks,
                "record": fig7_bc_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass for current Fig8 A/B-only source with no current access heatmap or colorbar",
    )

    fig7_layout_statuses, fig7_layout_record, fig7_layout_status_ok = step12_qc_record(
        "Fig8_layout_A_B_only_A_y_axis_recorded"
    )
    fig7_layout_record_checks = {
        "step12_qc_status_pass": fig7_layout_status_ok,
        "current_visual_panels_A_B": fig7_layout_record.get("current_visual_panels") == ["A", "B"],
        "former_panel_a_migrated_to_Fig6C": fig7_layout_record.get("former_panel_a_migrated_to") == "Fig6C",
        "panel_a_removed": fig7_layout_record.get("panel_a_removed") is True,
        "panel_a_heatmap_current_false": fig7_layout_record.get("panel_a_heatmap_current") is False,
        "panel_a_colorbar_current_false": fig7_layout_record.get("panel_a_colorbar_current") is False,
        "panel_b_c_side_by_side": fig7_layout_record.get("panel_b_c_bottom_side_by_side") is True,
        "panel_b_y_axis_restored": fig7_layout_record.get("panel_b_y_axis_restored") is True,
        "missing_source_fields_empty": fig7_layout_record.get("missing_source_fields", []) == [],
    }
    add(
        "Step12_QC_Fig8_layout_A_B_and_y_axis_synced",
        "pass" if all(fig7_layout_record_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig7_layout_statuses,
                "checks": fig7_layout_record_checks,
                "record": fig7_layout_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass for Fig8 A/B side-by-side layout with Panel A morphotype y-axis restored",
    )

    fig7_marker_statuses, fig7_marker_record, fig7_marker_status_ok = step12_qc_record(
        "Fig8_panel_C_marker_size_and_leader_lines_recorded"
    )
    fig7_leader_style = fig7_marker_record.get("leader_line_style_source", {})
    fig7_leader_labels = fig7_marker_record.get("leader_line_labels_source", [])
    fig7_marker_current = fig7_marker_record.get("marker_size_current_source_unique", [])
    fig7_marker_previous = fig7_marker_record.get("marker_size_previous_source_unique", [])
    fig7_marker_checks = {
        "step12_qc_status_pass": fig7_marker_status_ok,
        "marker_size_current_30": fig7_marker_current == [30.0],
        "marker_size_previous_42": fig7_marker_previous == [42.0],
        "leader_line_count_8": fig7_marker_record.get("leader_line_count_source") == 8,
        "leader_line_labels_match": fig7_leader_labels == expected_leader_labels,
        "leader_line_color": fig7_leader_style.get("color") == "#94A3B8",
        "leader_line_linewidth": fig7_leader_style.get("linewidth") == 0.45,
        "leader_line_alpha": fig7_leader_style.get("alpha") == 0.78,
        "missing_source_fields_empty": fig7_marker_record.get("missing_source_fields", []) == [],
    }
    add(
        "Step12_QC_Fig8_panel_B_marker_size_and_leader_lines_synced",
        "pass" if all(fig7_marker_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig7_marker_statuses,
                "checks": fig7_marker_checks,
                "record": fig7_marker_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass with Fig8 current Panel B marker size 42->30 points^2 and 8 leader lines for MT01, MT03, MT04, MT05, MT06, MT07, MT08, MT09 using #94A3B8, lw=0.45, alpha=0.78",
    )

    fig7_offset_statuses, fig7_offset_record, fig7_offset_status_ok = step12_qc_record(
        "Fig8_panel_C_MT03_MT04_label_adjustments_recorded"
    )
    fig7_expected_offsets = fig7_offset_record.get("expected_offsets", {})
    fig7_offset_revisions = fig7_offset_record.get("expected_offset_revisions", {})
    fig7_offset_checks = {
        "step12_qc_status_pass": fig7_offset_status_ok,
        "MT03_current_offset": fig7_expected_offsets.get("MT03") == [9, 5, "left"],
        "MT04_current_offset": fig7_expected_offsets.get("MT04") == [-11, -3, "right"],
        "MT03_revision_current": fig7_offset_revisions.get("MT03", {}).get("current") == [9, 5, "left"],
        "MT04_revision_current": fig7_offset_revisions.get("MT04", {}).get("current") == [-11, -3, "right"],
        "missing_source_fields_empty": fig7_offset_record.get("missing_source_fields", []) == [],
    }
    add(
        "Step12_QC_Fig8_MT03_MT04_label_adjustments_synced",
        "pass" if all(fig7_offset_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig7_offset_statuses,
                "checks": fig7_offset_checks,
                "record": fig7_offset_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass for Fig8 current Panel B MT03 [9,5,left] and MT04 [-11,-3,right] label-offset revisions",
    )

    fig7_vector_statuses, fig7_vector_record, fig7_vector_status_ok = step12_qc_record(
        "Fig8_vector_run_start_protection"
    )
    fig7_vector_reference_statuses, fig7_vector_reference_record, _ = step12_qc_record(
        "Fig8_Step13_material_reference_vector_status"
    )
    fig7_vector_reference_details = step12_qc_details("Fig8_Step13_material_reference_vector_status")
    fig7_vector_reference_matches = fig7_vector_reference_record.get("accepted_reference_matches", {})
    fig7_all_reference_hashes_match = all(
        bool(fig7_vector_reference_matches.get(key))
        for key in ["Fig8_png", "Fig8_svg", "Fig8_pdf", "Fig8_csv"]
    )
    fig7_clean_reference_ok = (
        fig7_vector_reference_statuses == ["pass"]
        and fig7_vector_reference_record.get("drift_scope") == "none"
        and fig7_all_reference_hashes_match
        and fig7_vector_reference_record.get("accepted_vector_mismatches", []) == []
        and fig7_vector_reference_record.get("accepted_support_mismatches", []) == []
        and fig7_vector_reference_record.get("missing", []) == []
    )
    fig7_recorded_vector_drift_ok = (
        fig7_vector_reference_statuses == ["warn"]
        and fig7_vector_reference_record.get("non_target_vector_metadata_drift_accepted") is True
        and fig7_vector_reference_record.get("drift_scope") == "vector_only"
        and fig7_vector_reference_record.get("accepted_vector_mismatches") == ["Fig8_svg"]
        and fig7_vector_reference_record.get("accepted_support_mismatches", []) == []
        and all(bool(fig7_vector_reference_matches.get(key)) for key in ["Fig8_png", "Fig8_pdf", "Fig8_csv"])
        and fig7_vector_reference_record.get("missing", []) == []
        and any(
            "non-target Matplotlib vector metadata/id drift" in detail
            and "run-start Fig8 vector bytes were restored" in detail
            for detail in fig7_vector_reference_details
        )
    )
    fig7_reference_unavailable_but_current_material_synced = (
        fig7_vector_reference_statuses == ["warn"]
        and set(fig7_vector_reference_record.get("missing", [])) == {"Fig8_png", "Fig8_svg", "Fig8_pdf", "Fig8_csv"}
        and fig7_vector_reference_record.get("mismatches", []) == []
        and bool(all_fig_sync_ok)
    )
    fig7_stale_reference_but_current_material_synced = (
        fig7_vector_reference_statuses == ["warn"]
        and bool(all_fig_sync_ok)
        and bool(fig7_vector_reference_record.get("mismatches", []))
    )
    fig7_vector_checks = {
        "run_start_guard_status_pass": fig7_vector_status_ok,
        "run_start_preserved": fig7_vector_record.get("preserved_from_run_start") is True,
        "run_start_final_mismatches_empty": fig7_vector_record.get("final_mismatches", []) == [],
        "run_start_protected_file_count_0": fig7_vector_record.get("protected_file_count") == 0,
        "run_start_restored_none": sorted(fig7_vector_record.get("restored", [])) == [],
        "material_reference_clean_or_recorded_vector_drift_or_current_step13_sync": fig7_clean_reference_ok
        or fig7_recorded_vector_drift_ok
        or fig7_reference_unavailable_but_current_material_synced
        or fig7_stale_reference_but_current_material_synced,
        "current_step13_material_hashes_synced": bool(all_fig_sync_ok),
        "recorded_vector_drift_not_needed_or_ok": fig7_clean_reference_ok
        or fig7_recorded_vector_drift_ok
        or fig7_reference_unavailable_but_current_material_synced
        or fig7_stale_reference_but_current_material_synced,
    }
    add(
        "Step12_QC_Fig8_vector_drift_caveat_and_guard_synced",
        "pass" if all(fig7_vector_checks.values()) else "fail",
        json.dumps(
            {
                "run_start_statuses": fig7_vector_statuses,
                "material_reference_statuses": fig7_vector_reference_statuses,
                "material_reference_details": fig7_vector_reference_details,
                "checks": fig7_vector_checks,
                "run_start_record": fig7_vector_record,
                "material_reference_record": fig7_vector_reference_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC verifies current Fig8 target files are allowed to update and Step13 material hashes are synced.",
    )

    fig1_protect_statuses, fig1_protect_record, fig1_protect_status_ok = step12_qc_record(
        "Fig1_run_start_generation_or_protection"
    )
    fig1_protect_checks = {
        "step12_qc_status_pass": fig1_protect_status_ok,
        "preserved_from_run_start": fig1_protect_record.get("preserved_from_run_start") is True,
        "final_mismatches_empty": fig1_protect_record.get("final_mismatches", []) == [],
        "protected_file_count_4": fig1_protect_record.get("protected_file_count") == 4,
    }
    add(
        "Step12_QC_Fig1_run_start_protection_synced",
        "pass" if all(fig1_protect_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig1_protect_statuses,
                "checks": fig1_protect_checks,
                "record": fig1_protect_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass for Fig1 PNG/SVG/PDF/source CSV run-start protection or first-run generation.",
    )

    fig1_reference_statuses, fig1_reference_record, fig1_reference_status_ok = step12_qc_record(
        "Fig1_Step13_material_reference"
    )
    fig1_reference_matches = fig1_reference_record.get("accepted_reference_matches", {})
    fig1_reference_available = fig1_reference_record.get("accepted_reference_available", {})
    fig1_reference_unavailable_but_current_material_synced = (
        fig1_reference_statuses == ["warn"]
        and fig1_reference_record.get("mismatches", []) == []
        and set(fig1_reference_record.get("missing", [])) == {"Fig1_png", "Fig1_svg", "Fig1_pdf", "Fig1_csv"}
        and bool(all_fig_sync_ok)
    )
    fig1_reference_checks = {
        "step12_qc_status_pass_or_current_step13_sync": fig1_reference_status_ok
        or fig1_reference_unavailable_but_current_material_synced,
        "all_reference_hashes_available_or_current_step13_sync": (
            bool(fig1_reference_available) and all(bool(v) for v in fig1_reference_available.values())
        )
        or fig1_reference_unavailable_but_current_material_synced,
        "all_reference_hashes_match_or_current_step13_sync": (
            bool(fig1_reference_matches) and all(bool(v) for v in fig1_reference_matches.values())
        )
        or fig1_reference_unavailable_but_current_material_synced,
        "mismatches_empty": fig1_reference_record.get("mismatches", []) == [],
        "missing_empty_or_current_step13_sync": fig1_reference_record.get("missing", []) == []
        or fig1_reference_unavailable_but_current_material_synced,
        "current_step13_material_hashes_synced": bool(all_fig_sync_ok),
    }
    add(
        "Step12_QC_Fig1_Step13_material_reference_synced",
        "pass" if all(fig1_reference_checks.values()) else "fail",
        json.dumps(
            {
                "step12_qc_statuses": fig1_reference_statuses,
                "checks": fig1_reference_checks,
                "record": fig1_reference_record,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 QC pass for Fig1 accepted Step13 material-reference hashes or current Step13 material sync",
    )

    fig8 = tables["fig10_city_inequality_step12"]
    fig8_required = {
        "city_id",
        "category",
        "population_weighted_access_share",
        "accessibility_gini",
        "no_access_population_share",
        "fig11_row_definition",
    }
    fig8_missing = sorted(fig8_required - set(fig8.columns))
    fig8_ok = (
        len(fig8) == 240
        and not fig8_missing
        and fig8["city_id"].nunique() == 48
        and fig8["category"].nunique() == 5
    )
    add(
        "Fig10_step12_city_accessibility_inequality_240_rows",
        "pass" if fig8_ok else "fail",
        json.dumps(
            {
                "rows": len(fig8),
                "city_count": int(fig8["city_id"].nunique()) if "city_id" in fig8.columns else 0,
                "category_count": int(fig8["category"].nunique()) if "category" in fig8.columns else 0,
                "missing_fields": fig8_missing,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "240 Step12 Fig10 city-level rows: 48 cities x 5 facility categories",
    )

    fig10 = tables["fig5_radar_step12"]
    fig10_ok = (
        len(fig10) == 70
        and fig10.get("morphotype", pd.Series(dtype=str)).nunique() == 10
        and fig10.get("score_metric", pd.Series(dtype=str)).nunique() == 7
    )
    add(
        "Fig5_radar_step12_rows",
        "pass" if fig10_ok else "fail",
        json.dumps({"rows": len(fig10), "morphotypes": int(fig10.get("morphotype", pd.Series(dtype=str)).nunique()), "score_metrics": int(fig10.get("score_metric", pd.Series(dtype=str)).nunique())}, ensure_ascii=False, sort_keys=True),
        "70 Step12 Fig5 radar rows: 10 morphotypes x 7 score metrics",
    )

    fig11 = tables["fig9_observed_step12"]
    fig11_policy = fig11.get("label_policy", pd.Series(dtype=str)).astype(str)
    fig11_ok = (
        len(fig11) == 30
        and fig11.get("morphotype", pd.Series(dtype=str)).nunique() == 10
        and fig11.get("outcome", pd.Series(dtype=str)).nunique() == 3
        and {"observed_weighted_group_mean", "weight_sum", "label_policy"}.issubset(fig11.columns)
        and fig11_policy.str.contains("observed weighted group means", case=False, regex=False, na=False).all()
        and fig11_policy.str.contains("not adjusted marginal effects", case=False, regex=False, na=False).all()
    )
    add(
        "Fig9_observed_weighted_group_means_step12_rows",
        "pass" if fig11_ok else "fail",
        json.dumps(
            {
                "rows": len(fig11),
                "morphotypes": int(fig11.get("morphotype", pd.Series(dtype=str)).nunique()),
                "outcomes": int(fig11.get("outcome", pd.Series(dtype=str)).nunique()),
                "label_policy_values": sorted(fig11_policy.dropna().unique().tolist()),
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "30 Step12 Fig9 rows: observed weighted group means / group marginal summary, not adjusted marginal effects",
    )

    fig7_core = tables["fig7_core_city_step12"]
    fig7_core_required = {
        "city_id",
        "city_name_en",
        "region",
        "quality_tier",
        "scale",
        "dominant_morphotype",
        "dominant_share",
        "morphotype_entropy",
        "type_changed_between_scales",
        "share_delta",
        "access_metric",
        "circuity_metric",
        "inequality_metric",
        "quality_score",
        "entropy_norm",
        "access_norm",
        "circuity_norm",
        "inequality_norm",
        "quality_norm",
        "sorting_rank",
        "dominant_share_source",
        "source_fields_fallback_notes",
    }
    fig7_core_missing = sorted(fig7_core_required - set(fig7_core.columns))
    fig7_core_scale_counts = fig7_core.get("scale", pd.Series(dtype=str)).astype(str).value_counts().to_dict()
    fig7_core_city_count = int(fig7_core.get("city_id", pd.Series(dtype=str)).nunique())
    if {"city_id", "scale"}.issubset(fig7_core.columns):
        fig7_core_city_scale_sets = fig7_core.groupby("city_id")["scale"].agg(lambda s: sorted(set(s.astype(str)))).to_dict()
        fig7_core_bad_scale_cities = {
            city_id: scales for city_id, scales in fig7_core_city_scale_sets.items() if scales != ["hex_1km", "hex_2km"]
        }
    else:
        fig7_core_bad_scale_cities = {}
    fig7_core_metric_cols = ["entropy_norm", "access_norm", "circuity_norm", "inequality_norm", "quality_norm"]
    fig7_core_metric_nonnull = {
        col: int(pd.to_numeric(fig7_core.get(col, pd.Series(dtype=float)), errors="coerce").notna().sum())
        for col in fig7_core_metric_cols
    }
    fig7_core_ok = (
        len(fig7_core) == 96
        and fig7_core_city_count == 48
        and fig7_core_scale_counts.get("hex_1km", 0) == 48
        and fig7_core_scale_counts.get("hex_2km", 0) == 48
        and not fig7_core_missing
        and not fig7_core_bad_scale_cities
        and all(v >= 80 for v in fig7_core_metric_nonnull.values())
    )
    add(
        "Fig7_core_city_overview_step12_rows",
        "pass" if fig7_core_ok else "fail",
        json.dumps(
            {
                "rows": len(fig7_core),
                "city_count": fig7_core_city_count,
                "scale_counts": fig7_core_scale_counts,
                "missing_fields": fig7_core_missing,
                "bad_scale_cities": fig7_core_bad_scale_cities,
                "metric_nonnull_rows": fig7_core_metric_nonnull,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "96 Step12 Fig7 rows: 48 core cities x hex_1km/hex_2km, with non-null normalized strip metrics",
    )

    fig11_model = tables["fig11_model_coeff_step12"]
    fig11_model_forbidden = fig11_model[
        fig11_model.get("term", pd.Series("", index=fig11_model.index)).astype(str).str.contains(r"Intercept|C\(city_id\)|fixed_effect|city FE", case=False, regex=True, na=False)
        | fig11_model.get("term_kind", pd.Series("", index=fig11_model.index)).astype(str).str.contains("intercept|fixed_effect", case=False, regex=True, na=False)
    ]
    fig11_model_ok = (
        not fig11_model.empty
        and fig11_model_forbidden.empty
        and fig11_model.get("figure_id", pd.Series("Fig11", index=fig11_model.index)).astype(str).eq("Fig11").all()
    )
    add(
        "Fig11_model_coefficient_forest_step12_rows",
        "pass" if fig11_model_ok else "fail",
        json.dumps(
            {
                "rows": len(fig11_model),
                "forbidden_rows": len(fig11_model_forbidden),
                "figure_ids": sorted(fig11_model.get("figure_id", pd.Series(dtype=str)).dropna().astype(str).unique().tolist()),
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        "Step12 Fig11 model coefficient forest source data are nonempty, labeled Fig11 and exclude intercept/fixed-effect rows",
    )

    material_text = " ".join(material.astype(str).fillna("").agg(" ".join, axis=1).tolist()) if not material.empty else ""
    old_material_terms = [
        "Fig1_quality_aware_multiscale_typology_framework",
        "Fig2_global_sample_city_overview",
        "Fig4_multiscale_metric_curves",
        "Fig6_scale_transition_matrix",
        "Fig6_scale_network_composition_matrix",
        "Fig7_accessibility_detour_performance_by_morphotype",
        "Fig7_accessibility_by_morphotype_category",
        "Fig8_city_level_accessibility_inequality",
        "Fig8_circuity_efficiency_by_morphotype",
        "Fig9_model_coefficient_forest",
        "Fig10_morphotype_radar_signatures",
        "Fig11_observed_weighted_group_means",
        "Fig12_model_coefficient_forest",
        "Fig12_core_city_morphotype_stability_and_performance",
    ]
    old_material_terms.extend([f"{'S'}Fig1", f"{'S'}Fig2"])
    old_material_hits = [term for term in old_material_terms if term in material_text]
    add("old_step12_figure_names_absent_from_material_list", "pass" if not old_material_hits else "fail", old_material_hits, "No pre-reorder canonical figure names or legacy supplemental ids")

    city_ineq = tables["city_inequality"]
    add("city_inequality_240_rows", "pass" if len(city_ineq) == 240 else "fail", len(city_ineq), 240)

    coef_rows = len(tables["model_coefficients"])
    add("model_coefficients_933_rows", "pass" if coef_rows == 933 else "fail", coef_rows, "933")
    robust_rows = len(tables["robustness_results"])
    add("robustness_results_3084_rows", "pass" if robust_rows == 3084 else "fail", robust_rows, "3084")
    robust_summary_rows = len(tables["robustness_summary"])
    add("robustness_summary_10_rows", "pass" if robust_summary_rows == 10 else "fail", robust_summary_rows, "10")

    for label, df in [("Step11", tables["step11_qc"]), ("Step12", tables["step12_qc"])]:
        status_col = "status" if "status" in df.columns else None
        fail_count = int(df[status_col].astype(str).str.lower().eq("fail").sum()) if status_col else 0
        warn_count = int(df[status_col].astype(str).str.lower().isin(["warn", "warning"]).sum()) if status_col else 0
        add(f"{label}_QC_failures_absent", "pass" if fail_count == 0 else "fail", fail_count, 0)
        add(
            f"{label}_QC_warnings_nonblocking",
            "warn" if warn_count else "pass",
            warn_count,
            0,
            "Known Step11/Step12 warnings are retained as non-blocking caveats.",
        )

    add(
        "protected_step09_to_step12_directories_unchanged",
        "pass" if not protected_changed else "fail",
        len(protected_changed),
        0,
        "; ".join(protected_changed[:20]),
    )
    add("output_dir_guard_scope", "pass", safe_rel(paths.output_dir, paths.root), f"inside data/{STEP_NAME}")

    qc = pd.DataFrame(checks)
    return qc


def write_readme(paths: StepPaths, qc: pd.DataFrame, material: pd.DataFrame) -> None:
    status_counts = qc["status"].value_counts().to_dict()
    lines = [
        "# Step13 Export Paper Tables and Appendix Materials",
        "",
        "This directory is generated by `code_upload/13_export_paper_tables_appendices/step13_export_tables_appendices.py`. Step13 only exports, explains, records provenance, and builds the material manifest; it does not remodel, modify upstream data, or rewrite Step12 figures.",
        "",
        "## Run Command",
        "",
        "```bash",
        f"python code_upload/13_export_paper_tables_appendices/step13_export_tables_appendices.py --root {paths.root} --overwrite",
        "```",
        "",
        "## Main-text Tables",
        "",
        f"Step13 exports {len(MAIN_TABLE_FILES)} main-text CSV tables for Tables 2-9; Table 1 is maintained in the manuscript.",
        "",
        *[f"- `{name}`" for name in MAIN_TABLE_FILES],
        "",
        "Table 1 is maintained in `docs/manuscript.docx`; Step13 does not update the manuscript and only records that policy in the material manifest.",
        "",
        "## Appendix Support CSV",
        "",
        *[f"- `{name}`" for name in APPENDIX_TABLE_FILES],
        "",
        "## Appendix Excel",
        "",
        *[f"- `{name}`" for name in APPENDIX_FILES],
        "",
        "Appendix D includes the `morphotype_definitions`, `city_morphotype_composition`, and `representative_cities_by_type` sheets for MT01-MT10 local morphotype definitions and city composition semantics.",
        f"Appendix D and `city_morphotype_composition_full.csv` use this dominant-field policy: {DOMINANT_COMPATIBILITY_NOTE}",
        "",
        "Appendix E includes only aggregated or moderate-size data: the 5,030-row aggregated `OD_detour_metrics.parquet`, Fig8 A/B-only accessibility/detour performance source data, 240 Fig10 city-level accessibility inequality rows, Fig5 morphotype radar source data, Fig9 observed weighted group means source data, 96 Fig7 core-city overview rows, and low-confidence sensitivity data when present. The roughly four-million-row `facility_classification.parquet` and large `OD_detour_samples.parquet` inputs are listed in the material manifest but not embedded in Excel.",
        "",
        "Appendix F includes 933 Step11 model coefficient rows, the new Fig11 model coefficient forest source table, 3,084 robustness result rows, 10 robustness summary rows, and QC warnings. The Step11 small-group warning is recorded as a non-blocking warning.",
        "",
        "## Interpretation Notes",
        "",
        *[f"- {note}" for note in RISK_NOTES],
        f"- {DOMINANT_COMPATIBILITY_NOTE}",
        "",
        "## Paths and Provenance",
        "",
        f"- The output path guard only allows writes to `{paths.root / 'data' / STEP_NAME}` or its subdirectories.",
        "- `paper_materials_manifest.csv` / `.xlsx` record main-text tables, appendices, Step12 figures, Step12 figure data, key upstream inputs, hashes, mtimes, row/column counts, and sources.",
        f"- Current material manifest record count: {len(material)}",
        f"- QC status summary: {status_counts}",
    ]
    (paths.output_dir / "README.md").write_text("\n".join(lines), encoding="utf-8")


def package_versions() -> dict[str, str | None]:
    packages = ["pandas", "numpy", "pyarrow", "openpyxl"]
    out: dict[str, str | None] = {}
    for pkg in packages:
        try:
            out[pkg] = importlib_metadata.version(pkg)
        except importlib_metadata.PackageNotFoundError:
            out[pkg] = None
    return out


def git_head(root: Path) -> str | None:
    try:
        proc = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=root,
            check=False,
            capture_output=True,
            text=True,
            timeout=10,
        )
        if proc.returncode == 0:
            return proc.stdout.strip()
    except Exception:
        return None
    return None


def write_repro(
    paths: StepPaths,
    args: argparse.Namespace,
    input_report: dict[str, Any],
    qc: pd.DataFrame,
    protected_before_digest: str,
    protected_after_digest: str,
    protected_changed: list[str],
) -> None:
    script_path = Path(__file__).resolve()
    repro = {
        "step": STEP_NAME,
        "created_at": now_iso(),
        "root": str(paths.root),
        "output_dir": str(paths.output_dir),
        "command": f"python {safe_rel(script_path, paths.root)} --root {paths.root} --overwrite",
        "args": vars(args),
        "python": sys.version,
        "platform": platform.platform(),
        "package_versions": package_versions(),
        "git_head": git_head(paths.root),
        "script_hash_sha256": sha256_file(script_path),
        "input_report_json": "step13_input_acceptance_report.json",
        "required_missing_count": input_report.get("required_missing_count"),
        "qc_status_counts": qc["status"].value_counts().to_dict(),
        "protected_upstream_manifest_before_sha256": protected_before_digest,
        "protected_upstream_manifest_after_sha256": protected_after_digest,
        "protected_upstream_changed_paths": protected_changed,
        "risk_notes": RISK_NOTES,
    }
    write_json(paths.output_dir / "reproducibility_status_manifest.json", repro)


def load_tables(paths: StepPaths) -> dict[str, pd.DataFrame]:
    tables: dict[str, pd.DataFrame] = {}
    tables["city_master"] = read_parquet(paths.city_master)
    tables["city_quality"] = read_csv(paths.city_quality)
    tables["type_centers"] = attach_english_labels(read_csv(paths.type_centers))
    tables["city_type_profiles"] = read_csv(paths.city_type_profiles)
    local_morphotype_cols = existing_parquet_columns(
        paths.local_morphotypes,
        ["network_type", "scale", "morphotype", "population_sum", "population_weight", "type_confidence"],
    )
    tables["local_morphotypes_for_totals"] = read_parquet(paths.local_morphotypes, columns=local_morphotype_cols)
    tables["od_metrics"] = read_parquet(paths.od_metrics)
    tables["fig7_accessibility_source"] = attach_english_labels(read_csv(paths.fig7_source))
    tables["fig7_detour_source"] = attach_english_labels(read_csv(paths.fig8_source))
    tables["fig5_radar_step12"] = attach_english_labels(read_csv(paths.step12_fig5_data))
    tables["fig6_scale_step12"] = attach_english_labels(read_csv(paths.step12_fig6_data))
    tables["fig7_core_city_step12"] = read_csv(paths.step12_fig7_data)
    tables["fig8_access_step12"] = attach_english_labels(read_csv(paths.step12_fig8_data))
    tables["fig9_observed_step12"] = attach_english_labels(read_csv(paths.step12_fig9_data))
    tables["fig10_city_inequality_step12"] = read_csv(paths.step12_fig10_data)
    tables["fig11_model_coeff_step12"] = read_csv(paths.step12_fig11_data)
    tables["city_inequality"] = read_csv(paths.city_inequality_source)
    tables["model_a"] = read_csv(paths.model_a_table)
    tables["model_results"] = read_csv(paths.model_results)
    tables["model_coefficients"] = read_csv(paths.model_coefficients)
    tables["robustness_results"] = read_csv(paths.robustness_results)
    tables["robustness_summary"] = read_csv(paths.robustness_summary)
    tables["marginal_effects"] = read_csv(paths.marginal_effects)
    tables["step11_qc"] = read_csv(paths.step11_qc)
    tables["step12_qc"] = read_csv(paths.step12_qc)
    if paths.quality_review.exists():
        tables["quality_review"] = read_csv(paths.quality_review)
    else:
        tables["quality_review"] = tables["city_quality"][tables["city_quality"].get("review_flag", False).fillna(False)].copy()
    if paths.type_contrib.exists():
        tables["type_contrib"] = read_csv(paths.type_contrib)
    else:
        tables["type_contrib"] = pd.DataFrame({"note": ["Feature contribution table not present."]})
    if paths.low_confidence_sensitivity.exists():
        tables["low_confidence_sensitivity"] = read_csv(paths.low_confidence_sensitivity)
    else:
        tables["low_confidence_sensitivity"] = pd.DataFrame({"note": ["low_confidence_sensitivity.csv not present."]})
    return tables


def build_and_write_outputs(paths: StepPaths, input_records: list[dict[str, Any]]) -> tuple[dict[str, pd.DataFrame], list[Path]]:
    tables = load_tables(paths)
    core_city_ids = set(tables["model_a"]["city_id"].dropna().astype(str).unique())
    model_b_city_ids = read_model_b_city_ids(paths.model_b_table)
    city_ineq_core_ids = set(tables["city_inequality"]["city_id"].dropna().astype(str).unique())
    core_city_ids = core_city_ids & model_b_city_ids & city_ineq_core_ids

    tables["table2"] = build_table2(tables["city_master"], core_city_ids)
    tables["table3"] = build_table3(tables["city_master"], core_city_ids)
    tables["table4"] = build_data_source_records(paths, input_meta_map(input_records))
    tables["table5"] = build_table5(paths)
    tables["table6"] = build_table6()
    tables["table7"] = build_table7(tables["model_results"])
    tables["table8"] = build_table8(tables["robustness_summary"])
    tables["city_morphotype_composition"] = build_city_morphotype_composition(
        tables["city_type_profiles"],
        tables["type_centers"],
    )
    tables["appendix_city_type_profiles"] = build_appendix_city_type_profiles(
        tables["city_type_profiles"],
        tables["type_centers"],
    )
    tables["table9"], tables["representative_cities_by_type"] = build_morphotype_definitions(
        tables["type_centers"],
        tables["fig5_radar_step12"],
        tables["city_morphotype_composition"],
        tables["local_morphotypes_for_totals"],
    )
    tables["core_cities"] = tables["city_master"][tables["city_master"]["city_id"].astype(str).isin(core_city_ids)].copy()
    tables["quality_warnings"] = build_quality_warning_sheet(paths)

    main_table_map = {
        "table2_main_sample_region_distribution.csv": tables["table2"],
        "table3_city_sample_stratification_logic.csv": tables["table3"],
        "table4_data_source_inventory.csv": tables["table4"],
        "table5_implementation_steps.csv": tables["table5"],
        "table6_analysis_modules.csv": tables["table6"],
        "table7_validation_model_design.csv": tables["table7"],
        "table8_robustness_checks.csv": tables["table8"],
        "table9_morphotype_definitions_and_representative_cities.csv": tables["table9"],
    }
    for name, df in main_table_map.items():
        write_csv(df, paths.output_dir / name)
    write_csv(tables["city_morphotype_composition"], paths.output_dir / "city_morphotype_composition_full.csv")

    city_master_quality = tables["city_master"].merge(
        tables["city_quality"],
        on=["city_id", "city_name_en", "iso3", "region", "sample_group"],
        how="left",
        suffixes=("", "_quality"),
    )
    quality_summary = (
        tables["city_quality"]
        .groupby(["region", "quality_tier"], dropna=False, as_index=False)
        .agg(
            city_count=("city_id", "nunique"),
            mean_quality_score=("quality_score", "mean"),
            mean_quality_weight=("quality_weight", "mean"),
            review_flag_count=("review_flag", "sum"),
        )
    )
    name_collision = (
        tables["type_centers"]
        .groupby("morphotype_name", dropna=False)
        .agg(name_count=("morphotype", "nunique"), morphotype_keys=("morphotype", lambda s: "; ".join(sorted(s.astype(str)))))
        .reset_index()
    )
    name_collision["unique_key_policy"] = "Use morphotype MT01-MT10 as the only type key; descriptive names may repeat."

    write_excel(
        paths.output_dir / "appendix_a_metric_dictionary.xlsx",
        {
            "variable_dictionary": pd.DataFrame(variable_dictionary_rows()),
            "source_table_columns": source_table_columns(paths),
            "main_table_fields": pd.DataFrame(
                [
                    {"main_table": name, "field_name": col}
                    for name, df in main_table_map.items()
                    for col in df.columns
                ]
            ),
        },
    )
    write_excel(
        paths.output_dir / "appendix_b_sample_city_inventory.xlsx",
        {
            "city_master_86": city_master_quality,
            "core_analysis_cities_48": tables["core_cities"],
            "region_summary": tables["table2"],
            "stratification_summary": tables["table3"],
        },
    )
    write_excel(
        paths.output_dir / "appendix_c_quality_score_details.xlsx",
        {
            "city_quality_scores": tables["city_quality"],
            "quality_summary": quality_summary,
            "quality_review_flags": tables["quality_review"],
            "qc_warnings": tables["quality_warnings"],
        },
    )
    write_excel(
        paths.output_dir / "appendix_d_clustering_and_type_naming.xlsx",
        {
            "morphotype_definitions": tables["table9"],
            "city_morphotype_composition": tables["city_morphotype_composition"],
            "representative_cities_by_type": tables["representative_cities_by_type"],
            "type_centers_MT_key": tables["type_centers"],
            "name_collision_check": name_collision,
            "city_type_profiles": tables["appendix_city_type_profiles"],
            "feature_contributions": tables["type_contrib"],
            "notes": build_appendix_notes("appendix_d", [RISK_NOTES[4], RISK_NOTES[5], DOMINANT_COMPATIBILITY_NOTE]),
        },
    )
    write_excel(
        paths.output_dir / "appendix_e_accessibility_detour_validation.xlsx",
        {
            "od_detour_metrics_5030": tables["od_metrics"],
            "fig8_access_detour_performance": tables["fig8_access_step12"],
            "fig10_city_inequality": tables["fig10_city_inequality_step12"],
            "fig5_morphotype_radar": tables["fig5_radar_step12"],
            "fig9_observed_group_means": tables["fig9_observed_step12"],
            "fig7_core_city_overview": tables["fig7_core_city_step12"],
            "low_conf_sensitivity": tables["low_confidence_sensitivity"],
            "notes": build_appendix_notes(
                "appendix_e",
                [
                    "OD_detour_metrics.parquet is an aggregated 5030-row table; OD_detour_samples.parquet is listed but not embedded.",
                    RISK_NOTES[1],
                    "Current Fig8 records only A/B panels: detour/circuity intervals and access-circuity quadrant assignments; the migrated access heatmap is now Fig6C.",
                    "Fig10 contains city-level accessibility inequality rows promoted from Step12 figure data.",
                    "Fig5 contains 10 morphotype x 7 score-metric radar source rows from Step12.",
                    "Fig9 contains observed weighted group means / group marginal summary only; it is not adjusted marginal effects.",
                    "Fig7 contains 96 long-format rows for the 48 core cities at hex_1km and hex_2km, plus city-level entropy/access/circuity/Gini/quality strip metrics and sorting metadata.",
                ],
            ),
        },
    )
    write_excel(
        paths.output_dir / "appendix_f_robustness_checks.xlsx",
        {
            "model_coefficients_933": tables["model_coefficients"],
            "fig11_model_coefficient_forest": tables["fig11_model_coeff_step12"],
            "robustness_results_3084": tables["robustness_results"],
            "robustness_summary_10": tables["robustness_summary"],
            "qc_warnings": tables["quality_warnings"],
            "notes": build_appendix_notes("appendix_f", [RISK_NOTES[0], RISK_NOTES[3]]),
        },
    )

    generated = [paths.output_dir / name for name in MAIN_TABLE_FILES + APPENDIX_TABLE_FILES + APPENDIX_FILES]
    return tables, generated


def run_step13(args: argparse.Namespace) -> dict[str, Any]:
    start = time.time()
    paths = get_paths(args.root, args.output_dir)
    paths.output_dir.mkdir(parents=True, exist_ok=True)
    check_existing_outputs(paths, args.overwrite)

    protected_dirs = [paths.step09_dir, paths.step10_dir, paths.step11_dir, paths.step12_dir]
    protected_before = collect_dir_manifest(paths.root, protected_dirs)
    protected_before_digest = manifest_digest(protected_before)

    input_records = collect_input_records(paths)
    input_report = write_input_report(paths, input_records, protected_before_digest)
    if input_report["required_missing_count"]:
        missing = ", ".join(input_report["required_missing"][:10])
        raise FileNotFoundError(f"Missing required Step13 inputs: {missing}")

    tables, generated = build_and_write_outputs(paths, input_records)

    protected_after = collect_dir_manifest(paths.root, protected_dirs)
    protected_after_digest = manifest_digest(protected_after)
    protected_changed = compare_manifests(protected_before, protected_after)

    material_paths = generated + [
        paths.output_dir / "paper_materials_manifest.csv",
        paths.output_dir / "paper_materials_manifest.xlsx",
        paths.output_dir / "step13_input_acceptance_report.json",
        paths.output_dir / "step13_input_acceptance_report.md",
        paths.output_dir / "step13_quality_checks.csv",
        paths.output_dir / "step13_run_log.json",
        paths.output_dir / "reproducibility_status_manifest.json",
        paths.output_dir / "README.md",
    ]
    material = build_material_rows(paths, material_paths, input_records)
    write_material_list(paths, material)

    qc = build_qc(paths, tables, protected_changed, material)
    write_csv(qc, paths.output_dir / "step13_quality_checks.csv")
    write_readme(paths, qc, material)
    write_repro(paths, args, input_report, qc, protected_before_digest, protected_after_digest, protected_changed)

    # Refresh material list after QC/README/repro exist; self hashes remain intentionally blank.
    material = build_material_rows(paths, material_paths, input_records)
    write_material_list(paths, material)

    run_log = {
        "step": STEP_NAME,
        "started_at": datetime.fromtimestamp(start).astimezone().isoformat(timespec="seconds"),
        "finished_at": now_iso(),
        "elapsed_seconds": round(time.time() - start, 3),
        "root": str(paths.root),
        "output_dir": str(paths.output_dir),
        "args": vars(args),
        "main_table_count": len(MAIN_TABLE_FILES),
        "appendix_count": len(APPENDIX_FILES),
        "record_file_count": len(RECORD_FILES),
        "qc_status_counts": qc["status"].value_counts().to_dict(),
        "qc_fail_count": int(qc["status"].eq("fail").sum()),
        "qc_warn_count": int(qc["status"].eq("warn").sum()),
        "protected_upstream_manifest_before_sha256": protected_before_digest,
        "protected_upstream_manifest_after_sha256": protected_after_digest,
        "protected_upstream_changed_paths": protected_changed,
        "generated_files": [safe_rel(p, paths.root) for p in material_paths],
        "risk_notes": RISK_NOTES,
    }
    write_json(paths.output_dir / "step13_run_log.json", run_log)

    # Final inventory and QC refresh happen after every required record file
    # exists, so the required-output check reflects the final package state.
    material = build_material_rows(paths, material_paths, input_records)
    write_material_list(paths, material)
    qc = build_qc(paths, tables, protected_changed, material)
    write_csv(qc, paths.output_dir / "step13_quality_checks.csv")
    write_readme(paths, qc, material)
    write_repro(paths, args, input_report, qc, protected_before_digest, protected_after_digest, protected_changed)
    run_log.update(
        {
            "finished_at": now_iso(),
            "elapsed_seconds": round(time.time() - start, 3),
            "qc_status_counts": qc["status"].value_counts().to_dict(),
            "qc_fail_count": int(qc["status"].eq("fail").sum()),
            "qc_warn_count": int(qc["status"].eq("warn").sum()),
        }
    )
    write_json(paths.output_dir / "step13_run_log.json", run_log)
    material = build_material_rows(paths, material_paths, input_records)
    write_material_list(paths, material)

    if int(qc["status"].eq("fail").sum()) > 0:
        raise RuntimeError("Step13 QC has fail rows. See step13_quality_checks.csv.")
    return run_log


def main() -> None:
    args = parse_args()
    run_log = run_step13(args)
    print(json.dumps(json_ready({
        "status": "ok",
        "output_dir": run_log["output_dir"],
        "main_table_count": run_log["main_table_count"],
        "appendix_count": run_log["appendix_count"],
        "record_file_count": run_log["record_file_count"],
        "qc_status_counts": run_log["qc_status_counts"],
    }), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
